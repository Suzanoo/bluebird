"""Placement-only overlay contracts; not Desktop Excel acceptance."""
from dataclasses import replace
from datetime import date
from io import BytesIO
from unittest import TestCase
from unittest.mock import patch
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from engine.domain import Settings, prepare
from engine.workbook import render, apply_timescale_overlay
from test_f1 import schedule
from test_f1_chart_integrity import check_package, check_line_fills, A, C


class OverlayTests(TestCase):
    def test_bounds_follow_weekly_monthly_periods_and_body(self):
        for start, finish, count in [(date(2026,1,2),date(2026,1,2),1),
                                     (date(2026,1,30),date(2026,2,2),3),
                                     (date(2026,1,2),date(2027,3,1),60)]:
            for method in ('Equal','Duration'):
                base = schedule()
                s = replace(base, activities=tuple(replace(base.activities[0],
                    key=f'a{i}', start=start, finish=finish) for i in range(count)))
                cfg = Settings(method); prepared = prepare(s,cfg)
                raw = render(s,cfg,prepared); check_package(self,raw)
                wb = load_workbook(BytesIO(raw))
                for name in ('Main','Monthly'):
                    ws = wb[name]; ch = ws._charts[0]
                    periods = sum(ws.cell(4,c).value is not None for c in range(14,ws.max_column+1))
                    scurve = next(r for r in range(5,ws.max_row+1) if ws.cell(r,3).value=='Period Plan')
                    self.assertEqual((ch.anchor._from.col,ch.anchor._from.row),(13,4))
                    self.assertEqual((ch.anchor.to.col,ch.anchor.to.row),(13+periods,scurve-1))
                wb.close()

    def test_only_overlay_chart_geometry_and_background_change(self):
        s=schedule();cfg=Settings('Duration');prepared=prepare(s,cfg)
        def old_placement(ch,n,scurve): ch.anchor=f'B{scurve+6}'
        with patch('engine.workbook.apply_timescale_overlay',side_effect=old_placement):
            before=render(s,cfg,prepared)
        after=render(s,cfg,prepared)
        with ZipFile(BytesIO(before)) as a, ZipFile(BytesIO(after)) as b:
            for name in a.namelist():
                if name.startswith('xl/worksheets/') or name in ('xl/charts/chart1.xml','xl/drawings/drawing1.xml'):
                    self.assertEqual(a.read(name),b.read(name),name)
            for number in (2,3):
                old=ET.fromstring(a.read(f'xl/charts/chart{number}.xml'))
                new=ET.fromstring(b.read(f'xl/charts/chart{number}.xml'))
                check_line_fills(self,new)
                for node in (new,new.find(f'{{{C}}}chart/{{{C}}}plotArea')):
                    props=node.find(f'{{{C}}}spPr')
                    self.assertIsNotNone(props.find(f'{{{A}}}noFill'))
                    self.assertIsNotNone(props.find(f'{{{A}}}ln/{{{A}}}noFill'))
                    node.remove(props)
                self.assertEqual(ET.tostring(old),ET.tostring(new))

    def test_reapply_after_python_roundtrip_is_idempotent(self):
        s=schedule();cfg=Settings('Equal');raw=render(s,cfg,prepare(s,cfg))
        wb=load_workbook(BytesIO(raw))
        for name in ('Main','Monthly'):
            ws=wb[name]; n=sum(ws.cell(4,c).value is not None for c in range(14,ws.max_column+1))
            scurve=next(r for r in range(5,ws.max_row+1) if ws.cell(r,3).value=='Period Plan')
            for _ in range(2): apply_timescale_overlay(ws._charts[0],n,scurve)
        out=BytesIO();wb.save(out);wb.close()
        check_package(self,out.getvalue())
        with ZipFile(out) as z:
            for number in (2,3):
                r=ET.fromstring(z.read(f'xl/charts/chart{number}.xml'))
                self.assertIsNotNone(r.find(f'{{{C}}}spPr/{{{A}}}noFill'))
                self.assertIsNotNone(r.find(f'{{{C}}}chart/{{{C}}}plotArea/{{{C}}}spPr/{{{A}}}noFill'))
