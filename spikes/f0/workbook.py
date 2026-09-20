"""F0 renderer. Coordinate maps stay here, never in the domain model."""
from io import BytesIO
from dataclasses import asdict
import json
from openpyxl import Workbook
from openpyxl.styles import Protection, PatternFill, Font
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter as col
from model import bases, weeks, plan_profile

def render(schedule, method):
    basis = bases(schedule, method)
    periods = weeks(schedule)
    wb = Workbook(); main = wb.active; main.title = 'Main'
    amount = wb.create_sheet('Activity Amount'); meta = wb.create_sheet('_Metadata')
    wb.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)
    main.append(['F0 PROOF — initial weighting: '+method+'; actual inputs are weekly increments (0..1).'])
    main.append(['Internal ID','WBS','Activity ID','Activity Name','Row kind','Start','Finish','Duration hours','Basis','Weight','Complete'] + periods)
    amount.append(['Planned Budget input only. Blank = effective zero; no effect on initial weights.'])
    amount.append(['WBS','Activity ID','Activity Name','Amount','Internal ID','Effective Amount','Input state'])
    by_wbs = {w.key:w for w in schedule.wbs}
    children = {}
    for w in schedule.wbs: children.setdefault(w.parent, []).append(w)
    tasks = {}
    for i,a in enumerate(schedule.activities): tasks.setdefault(a.wbs, []).append((i,a))
    activity_rows, summary_rows = [], []
    def walk(parent=None, depth=0, path=''):
        start_rows = len(activity_rows)
        for w in children.get(parent, []):
            wp = w.code if schedule.source_system == 'MSP' else path + ('.' if path else '') + w.code
            mr, ar = main.max_row+1, amount.max_row+1
            main.append([w.key,wp,'',w.name,'WBS'])
            amount.append([f"=Main!B{mr}",'',f"=Main!D{mr}",None,w.key])
            main.row_dimensions[mr].outlineLevel = min(depth,7)
            amount.row_dimensions[ar].outlineLevel = min(depth,7)
            subset = walk(w.key, depth+1, wp)
            summary_rows.append((mr,subset))
        for i,a in tasks.get(parent, []):
            r, ar = main.max_row+1, amount.max_row+1
            main.append([a.key,path,a.source_id,a.name,'Plan',a.start,a.finish,a.duration_hours,basis[i],None,None] + plan_profile(a,periods))
            main.append([a.key,path,a.source_id,a.name,'Actual',a.start,a.finish,None,0,None,None])
            last = col(11+len(periods))
            for rr in (r,r+1):
                main.cell(rr,10,f'=I{r}/$I$1')
                main.cell(rr,11,f'=SUM(L{rr}:{last}{rr})')
                main.row_dimensions[rr].outlineLevel = min(depth,7)
            for c in range(12,12+len(periods)):
                cell=main.cell(r+1,c); cell.protection=Protection(locked=False)
                cell.number_format='0.0%'; cell.fill=PatternFill('solid',fgColor='DBEAFE')
            amount.append([f'=Main!B{r}',f'=Main!C{r}',f'=Main!D{r}',None,a.key,f'=IF(D{ar}="",0,D{ar})',f'=IF(D{ar}="","Missing",IF(D{ar}=0,"Zero","Populated"))'])
            amount.cell(ar,4).protection=Protection(locked=False)
            amount.cell(ar,4).fill=PatternFill('solid',fgColor='DBEAFE')
            amount.cell(ar,4).number_format='#,##0.00'
            amount.row_dimensions[ar].outlineLevel=min(depth,7)
            amount.conditional_formatting.add(f'D{ar}',FormulaRule(formula=[f'OR(D{ar}="",D{ar}=0)'],fill=PatternFill('solid',fgColor='FEF3C7')))
            activity_rows.append((r,ar))
        return activity_rows[start_rows:]
    walk()
    # Avoid 8192-character formula limit for large schedules.
    main['I1']=f'=SUM(I3:I{main.max_row})'
    for r,subset in summary_rows:
        # Fixed denominator across all descendant activities, not active weeks.
        for c in [11]+list(range(12,12+len(periods))):
            letter=col(c)
            if subset:
                lo,hi=subset[0][0],subset[-1][0]+1
                main.cell(r,c,f'=IFERROR(SUMPRODUCT(I{lo}:I{hi},{letter}{lo}:{letter}{hi})/SUM(I{lo}:I{hi}),0)')
    # Project Plan and Actual use contiguous plan-only basis; Actual is shifted one row.
    r=main.max_row+2; main.cell(r,4,'Project Plan');main.cell(r+1,4,'Project Actual')
    lo,hi=activity_rows[0][0],activity_rows[-1][0]
    for c in [11]+list(range(12,12+len(periods))):
        letter=col(c)
        main.cell(r,c,f'=SUMPRODUCT(I{lo}:I{hi},{letter}{lo}:{letter}{hi})/$I$1')
        # WBS rows can intervene: use zero in basis rows and actual aligned +1.
        main.cell(r+1,c,f'=SUMPRODUCT(I{lo}:I{hi},{letter}{lo+1}:{letter}{hi+1})/$I$1')
    dv=DataValidation(type='decimal',operator='greaterThanOrEqual',formula1=0,allow_blank=True);dv.error='Enter nonnegative Planned Budget';dv.showErrorMessage=True;amount.add_data_validation(dv)
    for _,ar in activity_rows:dv.add(amount.cell(ar,4))
    av=DataValidation(type='decimal',operator='between',formula1=0,formula2=1,allow_blank=True);av.showErrorMessage=True;main.add_data_validation(av)
    for mr,_ in activity_rows:av.add(f'L{mr+1}:{col(11+len(periods))}{mr+1}')
    main.conditional_formatting.add(f'K3:K{main.max_row}',FormulaRule(formula=['K3>1'],fill=PatternFill('solid',fgColor='FECACA')))
    meta.append(['schema','f0-proof-1']);meta.append(['method',method]);meta.append(['import_id',schedule.import_id]);meta.append(['source_hash',schedule.source_hash]);meta.append(['source_system',schedule.source_system]);meta.append(['source_project',schedule.source_project]);meta.append(['source_project_fields',json.dumps(schedule.source_project_fields,ensure_ascii=False)])
    meta.append(['kind','JSON (identity/provenance; not a trusted rebuild import contract)'])
    for w in schedule.wbs:meta.append(['wbs',json.dumps(asdict(w),ensure_ascii=False)])
    for a in schedule.activities:meta.append(['activity',json.dumps(asdict(a),default=str,ensure_ascii=False)])
    meta.sheet_state='veryHidden'
    for ws in (main,amount):
        ws.freeze_panes='E3' if ws==main else 'D3';ws.sheet_properties.outlinePr.summaryBelow=False
        ws.protection.sheet=True;ws.protection.formatRows=False
        for cell in ws[2]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1E3A5F')
        for key,width in [('A',20),('B',20),('C',24),('D',38)]:ws.column_dimensions[key].width=width
    for row in main.iter_rows(min_row=3,min_col=10):
        for cell in row:cell.number_format='0.0%'
    main.column_dimensions['A'].hidden=True;amount.column_dimensions['E'].hidden=True
    # Imported strings never become Excel formulas.
    for row in main.iter_rows(min_row=3,max_col=8):
        for cell in row:
            if isinstance(cell.value,str):cell.data_type='s'
    for row in meta:
        for cell in row:cell.data_type='s'
    out=BytesIO();wb.save(out)
    if out.tell()>4_000_000:raise ValueError('F0 output exceeds 4 MB response cap')
    return out.getvalue()
