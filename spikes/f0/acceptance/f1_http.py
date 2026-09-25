"""Read-only F1 smoke for local Services or an explicitly authorized Preview.

No raw XML or workbook is logged or retained. Tiny synthetic F0 fixtures reused.
"""
import io
import json
from pathlib import Path
import sys
import time
from zipfile import ZipFile

import httpx
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from test_runtime import fixture


def smoke(base):
    results = []
    with httpx.Client(base_url=base.rstrip('/'),timeout=65,trust_env=False) as client:
        for source in ('msp','p6'):
            for method in ('Equal','Duration'):
                raw=fixture(source,True)
                started=time.perf_counter()
                response=client.post('/api/progress/convert',params={'method':method,'cutoff':'Friday','distribution':'auto'},
                                     content=raw,headers={'Content-Type':'application/xml'})
                assert response.status_code==200,(response.status_code,response.text[:100])
                assert response.headers['content-type'].startswith('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                assert response.headers['cache-control']=='no-store'
                assert 'progress.xlsx' in response.headers['content-disposition']
                with ZipFile(io.BytesIO(response.content)) as z: assert z.testzip() is None
                w=load_workbook(io.BytesIO(response.content))
                assert {'Main','Monthly','Dashboard','Activity Amount'}.issubset(w.sheetnames)
                assert w['Main']['N4'].value.weekday()==4
                w.close()
                results.append({'source':source,'method':method,'status':200,'xml_bytes':len(raw),
                                'xlsx_bytes':len(response.content),'elapsed_s':round(time.perf_counter()-started,4)})
        for raw,query,status in [(b'<bad','method=Equal',400),(b'','method=Equal',400),
            (b'x','method=Amount',422),(fixture('msp',True),'method=Equal&cutoff=bad',422),
            (b'x'*4_000_001,'method=Equal',413)]:
            response=client.post('/api/progress/convert?'+query,content=raw,headers={'Content-Type':'application/xml'})
            assert response.status_code==status,(response.status_code,status)
    print(json.dumps({'conversions':results,'rejections':5,'scope':'HTTP only; not browser or Desktop Excel acceptance'},indent=2))


if __name__=='__main__': smoke(sys.argv[1])
