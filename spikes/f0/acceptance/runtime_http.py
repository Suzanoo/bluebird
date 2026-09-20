"""Read-only smoke against a running local/authorized preview; no files retained."""
import httpx
import io
import json
from pathlib import Path
import sys
import time
from urllib.error import HTTPError
from urllib.request import Request, urlopen
from zipfile import ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from test_runtime import fixture
from openpyxl import load_workbook


def smoke(base):
    results = []
    for source in ('msp', 'p6'):
        for method in ('Equal', 'Duration'):
            raw = fixture(source, method == 'Duration')
            started = time.perf_counter()
            request = Request(base.rstrip('/') + '/api/f0/convert?method=' + method,
                              data=raw, headers={'Content-Type': 'application/xml'})
            with urlopen(request, timeout=65) as response:
                data = response.read()
                assert response.status == 200
                assert response.headers['Content-Type'].startswith('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                assert response.headers['Cache-Control'] == 'no-store'
                assert int(response.headers['Content-Length']) == len(data)
                assert 'progress-f0.xlsx' in response.headers['Content-Disposition']
                with ZipFile(io.BytesIO(data)) as z:
                    assert z.testzip() is None
                w = load_workbook(io.BytesIO(data))
                assert {'Main', 'Activity Amount'}.issubset(w.sheetnames)
                w.close()
                results.append(dict(source=source, method=method, status=200,
                                    xml_bytes=len(raw), xlsx_bytes=len(data),
                                    elapsed_s=round(time.perf_counter() - started, 4)))
    for raw, method, status in [(b'<broken', 'Equal', 400), (b'', 'Equal', 400),
                                (b'x', 'Amount', 422)]:
        try:
            urlopen(Request(base.rstrip('/') + '/api/f0/convert?method=' + method,
                            data=raw, headers={'Content-Type': 'application/xml'}), timeout=65)
        except HTTPError as e:
            assert e.code == status, (e.code, status)
            if status != 413:  # A platform cap may return non-JSON.
                assert 'code' in json.loads(e.read())['error']
        else:
            raise AssertionError('Expected request rejection')
    # httpx handles an early 413 while the request body is still being sent.
    r = httpx.post(base.rstrip('/') + '/api/f0/convert?method=Equal',
                   content=b'x' * 4_000_001, headers={'Content-Type': 'application/xml'},
                   timeout=65, trust_env=False)
    assert r.status_code == 413, r.status_code
    print(json.dumps({'conversions': results, 'rejection_cases': 4}, indent=2))


if __name__ == '__main__':
    smoke(sys.argv[1])
