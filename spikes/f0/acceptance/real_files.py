"""Acceptance runner: real files stay external to the repository.
Run: python spikes/f0/acceptance/real_files.py INPUT_DIR OUTPUT_DIR
"""
import sys,json,subprocess,collections,hashlib,io,re,functools
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from model import parse_xml,normalize,value,bases,weeks,plan_profile,aggregate
from openpyxl import load_workbook

def numeric_engine(ws):
    @functools.lru_cache(None)
    def n(e):
        if e is None:return 0.0
        if isinstance(e,(int,float)):return float(e)
        e=e.lstrip('=').replace('$','')
        try:return float(e)
        except ValueError:pass
        if e.startswith('IFERROR('):
            try:return n(e[8:-3])
            except ZeroDivisionError:return 0.0
        if '/' in e:
            l,r=e.split('/',1);return n(l)/n(r)
        if e.startswith('SUMPRODUCT('):
            a,b=e[11:-1].split(',');return sum(n(x.value)*n(y.value) for x,y in zip(cells(a),cells(b)))
        if e.startswith('SUM('):return sum(n(c.value) for c in cells(e[4:-1]))
        if re.fullmatch('[A-Z]+[0-9]+',e):return n(ws[e].value)
        raise AssertionError('Unsupported acceptance formula: '+e)
    def cells(r):return [c for row in ws[r] for c in row]
    return n

def close(a,b):assert abs(a-b)<1e-9,(a,b)

def main(input_dir,out):
    out.mkdir(parents=True,exist_ok=True)
    sources={k:(input_dir/('009_'+k+'.xml')).read_bytes() for k in ('MSP','P6')}
    roots={k:parse_xml(v) for k,v in sources.items()}
    p6=normalize(roots['P6'],sources['P6']);pids={a.source_id for a in p6.activities}
    leaves=[a for a in roots['MSP'].findall('Tasks/Task') if value(a,'Summary')!='1']
    candidates=collections.defaultdict(list)
    for a in leaves:
        for e in a.findall('ExtendedAttribute'):candidates[value(e,'FieldID')].append(value(e,'Value'))
    fields=[k for k,v in candidates.items() if len(v)==len(leaves) and len(set(v))==len(v) and set(v)==pids]
    assert len(fields)==1,'No unique paired-file Activity ID correspondence'
    mapping=fields[0]
    msp=normalize(roots['MSP'],sources['MSP'],msp_activity_id_field=mapping)
    models={'MSP':msp,'P6':p6};pmap={a.source_id:a for a in p6.activities}
    def hierarchy(s,k,attr):
        by={w.key:w for w in s.wbs};nodes=[]
        while k:
            w=by[k];nodes.append(getattr(w,attr));k=w.parent
        return tuple(reversed(nodes))
    diffs=collections.Counter();detail=[]
    for a in msp.activities:
        b=pmap[a.source_id];d=[]
        for attr in ('name','start','finish','duration_hours','milestone','calendar_id'):
            if getattr(a,attr)!=getattr(b,attr):diffs[attr]+=1;d.append(attr)
        if hierarchy(msp,a.wbs,'name')!=hierarchy(p6,b.wbs,'name'):diffs['wbs_name_path']+=1;d.append('wbs_name_path')
        if d:detail.append({'id':a.source_id,'fields':d})
    result={'mapping':{'field':mapping,'definition_count':len(roots['MSP'].findall('ExtendedAttributes/ExtendedAttribute')),'basis':'Unique complete value-set match to supplied P6 IDs; pair-specific evidence, not XML-declared semantic alias'},'comparison':{'matched':len(pids),'differences':dict(diffs),'details':detail},'files':{},'runs':[]}
    # Measure every conversion before parent-side workbook inspection grows RSS.
    measured = {}
    for kind in ('MSP','P6'):
        for method in ('Equal','Duration'):
            dest = out/f'009_{kind}_{method}_F0.xlsx'
            cmd=[sys.executable,str(ROOT/'run.py'),str(input_dir/f'009_{kind}.xml'),str(dest),'--method',method]
            if kind=='MSP':cmd+=['--msp-activity-id-field',mapping]
            measured[(kind,method)]=json.loads(subprocess.check_output(cmd,text=True))
    import html
    result['comparison']['names_equal_after_one_html_decode'] = all(html.unescape(a.name)==pmap[a.source_id].name for a in msp.activities)
    def descendants(s):
        by={w.key:set() for w in s.wbs};parents={w.key:w.parent for w in s.wbs}
        for a in s.activities:
            k=a.wbs
            while k:by[k].add(a.source_id);k=parents[k]
        return collections.Counter(tuple(sorted(v)) for v in by.values())
    result['comparison']['wbs_descendant_sets_equal'] = descendants(msp)==descendants(p6)
    result['comparison']['wbs_paths_equal_after_one_html_decode'] = all(tuple(html.unescape(v) for v in hierarchy(msp,a.wbs,'name'))==hierarchy(p6,pmap[a.source_id].wbs,'name') for a in msp.activities)
    for kind,s in models.items():
        r=roots[kind];p=r if kind=='MSP' else r.find('Project');activities=leaves if kind=='MSP' else p.findall('Activity')
        rawkeys=['DurationFormat','CalendarUID','Milestone','Summary','Manual','PercentComplete'] if kind=='MSP' else ['CalendarObjectId','Type','Status','PercentCompleteType']
        result['files'][kind]={'sha256':hashlib.sha256(sources[kind]).hexdigest(),'xml_bytes':len(sources[kind]),'project':s.source_project_fields,'activity_count':len(s.activities),'wbs_count':len(s.wbs),'weeks':len(weeks(s)),'start':str(min(a.start for a in s.activities)),'finish':str(max(a.finish for a in s.activities)),'duration_hours':sum(bases(s,'Duration')),'duration_missing':sum(a.duration_hours is None for a in s.activities),'milestones':[a.source_id for a in s.activities if a.milestone],'max_wbs_depth':max(len(hierarchy(s,w.key,'name')) for w in s.wbs),'duplicate_ids':len(s.activities)-len({a.source_id for a in s.activities}),'duplicate_objects':len(s.activities)-len({a.source_object_id for a in s.activities}),'source_counts':{key:dict(collections.Counter(value(a,key) for a in activities)) for key in rawkeys}}
        for method in ('Equal','Duration'):
            filename=f'009_{kind}_{method}_F0.xlsx';dest=out/filename
            cmd=[sys.executable,str(ROOT/'run.py'),str(input_dir/f'009_{kind}.xml'),str(dest),'--method',method]
            if kind=='MSP':cmd+=['--msp-activity-id-field',mapping]
            metrics=measured[(kind,method)];metrics.update(source=kind,file=filename,wbs=len(s.wbs),weeks=len(weeks(s)))
            w=load_workbook(dest);m=w['Main'];amt=w['Activity Amount'];assert w.sheetnames==['Main','Activity Amount','_Metadata']
            planrows={m.cell(row,3).value:row for row in range(3,m.max_row+1) if m.cell(row,5).value=='Plan'}
            amountrows={amt.cell(row,5).value:row for row in range(3,amt.max_row+1) if amt.cell(row,6).data_type=='f'}
            assert len(planrows)==len(amountrows)==len(s.activities)==208
            assert m.protection.sheet and amt.protection.sheet
            assert not m.sheet_properties.outlinePr.summaryBelow and not amt.sheet_properties.outlinePr.summaryBelow
            assert len(amt.conditional_formatting)==208
            for sid,row in planrows.items():
                assert m.cell(row,9).protection.locked and not m.cell(row+1,12).protection.locked
                ar=amountrows[m.cell(row,1).value];assert amt.cell(ar,4).value is None and not amt.cell(ar,4).protection.locked
                assert amt.cell(ar,2).value==f'=Main!C{row}' and amt.cell(ar,2).protection.locked
                assert amt.row_dimensions[ar].outlineLevel>0
            project_plan=next(row for row in range(3,m.max_row+1) if m.cell(row,4).value=='Project Plan');project_actual=project_plan+1
            evaluate=numeric_engine(m);close(evaluate(m.cell(project_plan,11).value),1);close(evaluate(m['I1'].value),sum(bases(s,method)))
            close(sum(evaluate(m.cell(row,10).value) for row in planrows.values()),1)
            for a in s.activities:close(sum(plan_profile(a,weeks(s))),1)
            # Exercise a known real Activity and preserved blank/zero/positive inputs.
            sid='A1000';pr=planrows[sid];ar=amountrows[m.cell(pr,1).value];other=[v for v in amountrows.values() if v!=ar]
            m.cell(pr+1,12,1);amt.cell(ar,4,0);amt.cell(other[0],4,1234.5)
            formula_before={c.coordinate:c.value for row in m for c in row if c.data_type=='f'}
            expected=(1/208 if method=='Equal' else next(a.duration_hours for a in s.activities if a.source_id==sid)/34464)
            evaluate=numeric_engine(m);close(evaluate(m.cell(project_actual,11).value),expected)
            buf=io.BytesIO();w.save(buf);again=load_workbook(io.BytesIO(buf.getvalue()));mm=again['Main'];aa=again['Activity Amount']
            assert aa.cell(ar,4).value==0 and aa.cell(other[0],4).value==1234.5 and aa.cell(other[1],4).value is None
            assert mm.cell(pr+1,12).value==1 and not mm.cell(pr+1,12).protection.locked
            assert formula_before=={c.coordinate:c.value for row in mm for c in row if c.data_type=='f'}
            assert again['_Metadata']['B2'].value==method and again['_Metadata'].sheet_state=='veryHidden'
            assert [x.outlineLevel for x in amt.row_dimensions.values()]==[x.outlineLevel for x in aa.row_dimensions.values()]
            close(numeric_engine(mm)(mm.cell(project_actual,11).value),expected)
            assert load_workbook(io.BytesIO(buf.getvalue()),data_only=True)['Main'].cell(project_actual,11).value is None
            # At 100% for every real activity, project actual reaches 100%.
            for row in planrows.values():mm.cell(row+1,12,1)
            close(numeric_engine(mm)(mm.cell(project_actual,11).value),1)
            metrics.update(automated_lifecycle='PASS (not Excel)',plan_total=1,all_complete_total=1,denominator=sum(bases(s,method)),manual={'activity':sid,'actual_cell':f'L{pr+1}','amount_zero_cell':f'D{ar}','amount_positive_cell':f'D{other[0]}','amount_blank_cell':f'D{other[1]}','project_plan':f'K{project_plan}','project_actual':f'K{project_actual}','expected_one_activity_percent':expected*100})
            result['runs'].append(metrics)
    (out/'real_file_evidence.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=='__main__':main(Path(sys.argv[1]).resolve(),Path(sys.argv[2]).resolve())
