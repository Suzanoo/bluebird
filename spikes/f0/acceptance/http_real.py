"""Real-file local HTTP proof. No hosted deployment claim."""
import io,json,sys,threading,tempfile,urllib.request
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from http_proof import Handler,HTTPServer
from openpyxl import load_workbook

def main(inputs):
    server=HTTPServer(('127.0.0.1',0),Handler);thread=threading.Thread(target=server.serve_forever);thread.start();results=[]
    old_tmp=tempfile.tempdir
    try:
        with tempfile.TemporaryDirectory() as td:
            tempfile.tempdir=td
            for source in ('MSP','P6'):
                for method in ('Equal','Duration'):
                    url=f'http://127.0.0.1:{server.server_port}/?method={method}'
                    if source=='MSP':url+='&msp_activity_id_field=188743731'
                    raw=(inputs/f'009_{source}.xml').read_bytes()
                    with urllib.request.urlopen(urllib.request.Request(url,data=raw,headers={'Content-Type':'application/xml'}),timeout=60) as res:
                        data=res.read();assert res.status==200 and res.headers['Cache-Control']=='no-store'
                        assert res.headers['Content-Length']==str(len(data))
                        w=load_workbook(io.BytesIO(data));assert w['_Metadata']['B2'].value==method
                        assert sum(w['Main'].cell(r,5).value=='Plan' for r in range(1,w['Main'].max_row+1))==208
                        assert not list(Path(td).iterdir()),'Leaked temporary worksheet file'
                        results.append({'source':source,'method':method,'http_status':res.status,'xlsx_bytes':len(data),'temporary_files_after_response':0})
    finally:
        tempfile.tempdir=old_tmp;server.shutdown();thread.join();server.server_close()
    print(json.dumps(results,indent=2))
if __name__=='__main__':main(Path(sys.argv[1]))
