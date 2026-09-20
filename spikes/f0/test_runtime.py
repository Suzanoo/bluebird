"""Transport regression; accepted XML/workbook semantics remain in test_proof."""
import asyncio
import io
import json
from pathlib import Path
import threading
import unittest
from unittest.mock import patch
from zipfile import ZipFile

import httpx
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.worksheet._writer import ALL_TEMP_FILES, create_temporary_file
from app import app, generate, MAX_BYTES, XLSX_TYPE

FIX = Path(__file__).parent / 'fixtures'


def fixture(source, duration=False):
    raw = (FIX / f'{source}_n8.xml').read_bytes()
    if duration:
        if source == 'msp':
            raw = raw.replace(b'<Summary>0</Summary>', b'<Summary>0</Summary><Duration>PT8H30M0S</Duration><DurationFormat>7</DurationFormat>')
        else:
            raw = raw.replace(b'<Id>A1290</Id>', b'<Id>A1290</Id><PlannedDuration>8.5</PlannedDuration>')
            raw = raw.replace(b'<Id>A1310</Id>', b'<Id>A1310</Id><PlannedDuration>16</PlannedDuration>')
    return raw


class RuntimeTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def post(self, raw=b'', method='Equal', **kwargs):
        return self.client.post('/api/f0/convert?method=' + method, content=raw,
                                headers={'Content-Type': 'application/xml'}, **kwargs)

    def test_sources_methods_headers_integrity(self):
        for source in ('msp', 'p6'):
            for method in ('Equal', 'Duration'):
                with self.subTest(source=source, method=method):
                    before = set(ALL_TEMP_FILES)
                    r = self.post(fixture(source, method == 'Duration'), method)
                    self.assertEqual(r.status_code, 200, r.text[:100] if r.status_code != 200 else '')
                    self.assertEqual(r.headers['content-type'], XLSX_TYPE)
                    self.assertEqual(r.headers['content-length'], str(len(r.content)))
                    self.assertEqual(r.headers['cache-control'], 'no-store')
                    self.assertEqual(r.headers['content-disposition'], 'attachment; filename="progress-f0.xlsx"')
                    with ZipFile(io.BytesIO(r.content)) as z:
                        self.assertIsNone(z.testzip())
                        self.assertIn('xl/workbook.xml', z.namelist())
                    w = load_workbook(io.BytesIO(r.content))
                    self.assertIn('Main', w.sheetnames)
                    self.assertIn('Activity Amount', w.sheetnames)
                    self.assertTrue(any(c.data_type == 'f' for row in w['Main'] for c in row))
                    w.close()
                    self.assertEqual(set(ALL_TEMP_FILES), before)

    def test_invalid_inputs(self):
        cases = [(b'', 'Equal', 400), (b'<bad', 'Equal', 400),
                 (b'<unknown/>', 'Equal', 422), (fixture('msp'), 'Amount', 422),
                 (fixture('msp'), '', 422), (fixture('msp'), 'Duration', 422),
                 (b'<!DOCTYPE x [<!ENTITY y "secret">]><x>&y;</x>', 'Equal', 422)]
        for raw, method, status in cases:
            with self.subTest(raw=raw[:20], method=method):
                r = self.post(raw, method)
                self.assertEqual(r.status_code, status)
                self.assertEqual(r.headers['cache-control'], 'no-store')
                self.assertIn('code', r.json()['error'])
        self.assertEqual(self.client.get('/api/f0/convert').status_code, 405)
        self.assertEqual(self.client.post('/api/f0/convert', content=b'x').status_code, 422)
        self.assertEqual(self.client.post('/api/f0/convert?method=Equal', content=b'x').status_code, 415)

    def test_declared_oversize(self):
        with patch('app.convert') as convert:
            r = self.post(b'x' * (MAX_BYTES + 1))
            self.assertEqual(r.status_code, 413)
            convert.assert_not_called()

    def test_streamed_oversize_without_length(self):
        async def check():
            async def body():
                yield b'x' * MAX_BYTES
                yield b'x'
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://test') as c:
                return await c.post('/api/f0/convert?method=Equal', content=body(), headers={'Content-Type': 'application/xml'})
        with patch('app.convert') as convert:
            self.assertEqual(asyncio.run(check()).status_code, 413)
            convert.assert_not_called()

    def test_conversion_failure_cleanup_preserves_unowned_file(self):
        unowned = create_temporary_file()
        owned = []
        def fail(filename, *args, **kwargs):
            owned.append(filename)
            raise OSError('PRIVATE XML sentinel')
        try:
            with patch('zipfile.ZipFile.write', side_effect=fail):
                r = self.post(fixture('msp'))
            self.assertEqual(r.status_code, 500)
            self.assertNotIn('PRIVATE', r.text)
            self.assertTrue(Path(unowned).exists())
            self.assertTrue(all(not Path(p).exists() for p in owned))
        finally:
            Path(unowned).unlink(missing_ok=True)
            ALL_TEMP_FILES.remove(unowned)

    def test_output_limit_and_late_result(self):
        with patch('app.convert', side_effect=ValueError('F0 output exceeds 4 MB response cap')):
            self.assertEqual(self.post(fixture('msp')).status_code, 413)
        with patch('app.convert', return_value=(b'x' * (MAX_BYTES + 1), {})):
            self.assertEqual(self.post(fixture('msp')).status_code, 413)
        with patch('app.convert', return_value=(b'PK', {})), patch('app.monotonic', side_effect=[0, 61]):
            self.assertEqual(self.post(fixture('msp')).status_code, 504)

    def test_concurrent_admission_and_recovery(self):
        entered, release = threading.Event(), threading.Event()
        def slow(*args):
            entered.set()
            release.wait(5)
            return b'PK', {}
        with patch('app.convert', side_effect=slow):
            worker = threading.Thread(target=generate, args=(b'x', 'Equal'))
            worker.start()
            try:
                self.assertTrue(entered.wait(5))
                self.assertEqual(generate(b'x', 'Equal').status_code, 503)
            finally:
                release.set(); worker.join(5)
        self.assertEqual(self.post(fixture('msp')).status_code, 200)

    def test_length_and_encoding(self):
        for headers, status in [({'Content-Length': 'abc'}, 400),
                                ({'Content-Length': '1'}, 400),
                                ({'Content-Encoding': 'gzip'}, 415)]:
            r = self.client.post('/api/f0/convert?method=Equal', content=b'<x/>',
                                 headers={'Content-Type': 'application/xml', **headers})
            self.assertEqual(r.status_code, status)


if __name__ == '__main__':
    unittest.main()
