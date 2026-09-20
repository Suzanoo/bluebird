import unittest,io,threading,urllib.request,urllib.error
from pathlib import Path
from dataclasses import replace
from openpyxl import load_workbook
from model import *
from workbook import render
from http_proof import Handler,HTTPServer
FIX=Path(__file__).parent/'fixtures'
def load(name='msp_n8.xml'):
    raw=(FIX/name).read_bytes();return normalize(parse_xml(raw),raw,msp_activity_id_field='188743731')
def duration_model():
    s=load();return replace(s,activities=tuple(replace(a,duration_hours=d) for a,d in zip(s.activities,[8,24])))
class Proof(unittest.TestCase):
    def test_equivalent_sources_and_preserved_identity(self):
        m,p=load(),load('p6_n8.xml')
        self.assertEqual([(a.source_id,a.start,a.finish) for a in m.activities],[(a.source_id,a.start,a.finish) for a in p.activities])
        self.assertEqual(m.activities[0].source_object_id,'4');self.assertIsNone(p.activities[0].source_object_id)
        self.assertEqual(len(m.wbs),4);self.assertNotEqual(m.activities[0].key,load().activities[0].key)
        raw=(FIX/'msp_n8.xml').read_bytes();self.assertEqual(m,normalize(parse_xml(raw),raw,m.import_id,msp_activity_id_field='188743731'))
    def test_equal_and_duration_oracle(self):
        s=duration_model();self.assertAlmostEqual(aggregate(bases(s,'Equal'),[1,0]),.5)
        self.assertAlmostEqual(aggregate(bases(s,'Duration'),[1,0]),.25)
        self.assertEqual(sum(bases(s,'Duration')),32)
        self.assertAlmostEqual(aggregate([8,24],[.5,.5]),.5)
    def test_no_default_or_amount(self):
        for choice in ('','Amount',None):
            with self.assertRaises(ValueError):bases(load(),choice)
    def test_missing_duration_no_silent_fallback(self):
        with self.assertRaises(ValueError):bases(load(),'Duration')
        self.assertEqual(bases(load(),'Equal'),[1,1])
    def test_milestones_zero_invalid_and_denominator(self):
        s=duration_model();a,b=s.activities
        s=replace(s,activities=(replace(a,duration_hours=0,milestone=True),b))
        self.assertEqual(bases(s,'Duration'),[0,24]);self.assertEqual(bases(s,'Equal'),[1,1])
        for pair in ((replace(a,duration_hours=0),b),(replace(a,duration_hours=None),b),(replace(a,duration_hours=0,milestone=True),replace(b,duration_hours=0,milestone=True))):
            with self.assertRaises(ValueError):bases(replace(s,activities=pair),'Duration')
    def test_duration_adapter_fields(self):
        for name,old,new in [('msp_n8.xml',b'<Summary>0</Summary>',b'<Summary>0</Summary><Duration>PT8H30M0S</Duration><DurationFormat>7</DurationFormat>'),('p6_n8.xml',b'<Id>A1290</Id>',b'<Id>A1290</Id><PlannedDuration>8.5</PlannedDuration>')]:
            raw=(FIX/name).read_bytes().replace(old,new);s=normalize(parse_xml(raw),raw,msp_activity_id_field='188743731');self.assertEqual(s.activities[0].duration_hours,8.5)
        raw=(FIX/'msp_n8.xml').read_bytes().replace(b'<Summary>0</Summary>',b'<Summary>0</Summary><Duration>PT8H</Duration><DurationFormat>8</DurationFormat>');self.assertIsNone(normalize(parse_xml(raw),raw,msp_activity_id_field='188743731').activities[0].duration_hours)
    def test_plan_conservation_and_fixed_denominator(self):
        s=load();self.assertTrue(all(abs(sum(plan_profile(a,weeks(s)))-1)<1e-12 for a in s.activities))
        self.assertEqual(aggregate([1,1],[1,0]),.5)
    def test_hostile_xml_and_structure(self):
        for raw in (b'<!DOCTYPE x [<!ENTITY x "a">]><Project>&x;</Project>',b'x'*4_000_001):
            with self.assertRaises(Exception):parse_xml(raw)
        raw=(FIX/'p6_n8.xml').read_bytes().replace(b'<Id>A1310</Id>',b'<Id>A1290</Id>')
        with self.assertRaises(ValueError):normalize(parse_xml(raw),raw,msp_activity_id_field='188743731')
        raw=(FIX/'p6_n8.xml').read_bytes().replace(b'<WBSObjectId>30337</WBSObjectId>',b'<WBSObjectId>999</WBSObjectId>')
        with self.assertRaises(ValueError):normalize(parse_xml(raw),raw,msp_activity_id_field='188743731')
    def test_workbook_contract_and_roundtrip(self):
        s=duration_model();data=render(s,'Duration');w=load_workbook(io.BytesIO(data));m,a=w['Main'],w['Activity Amount']
        self.assertEqual(w.calculation.calcMode,'auto');self.assertEqual(w['_Metadata'].sheet_state,'veryHidden')
        ar=[r for r in range(3,a.max_row+1) if a.cell(r,6).data_type=='f'];self.assertEqual(len(ar),2)
        self.assertFalse(a.cell(ar[0],4).protection.locked);self.assertTrue(a.cell(ar[0],2).protection.locked)
        self.assertGreater(a.row_dimensions[ar[0]].outlineLevel,0);self.assertEqual(len(a.conditional_formatting),2)
        a.cell(ar[0],4,0);a.cell(ar[1],4,1234.50)
        actual=next(r for r in range(3,m.max_row+1) if m.cell(r,5).value=='Actual');m.cell(actual,12,.25)
        buf=io.BytesIO();w.save(buf);again=load_workbook(io.BytesIO(buf.getvalue()))
        self.assertEqual(again['Activity Amount'].cell(ar[0],4).value,0);self.assertEqual(again['Main'].cell(actual,12).value,.25)
        self.assertEqual(again['Main'].cell(actual,11).data_type,'f')
        # Missing caches are expected: openpyxl is not an Excel engine.
        cached=load_workbook(io.BytesIO(buf.getvalue()),data_only=True);self.assertIsNone(cached['Main'].cell(actual,11).value)
    def test_formula_injection_is_text(self):
        s=load();s=replace(s,activities=(replace(s.activities[0],name='=HYPERLINK("https://example.invalid")'),s.activities[1]))
        w=load_workbook(io.BytesIO(render(s,'Equal')))
        cells=[c for row in w['Main'] for c in row if c.value==s.activities[0].name];self.assertTrue(cells);self.assertTrue(all(c.data_type=='s' for c in cells))
    def test_generated_main_formulas_against_independent_oracle(self):
        # Restricted evaluator for this renderer's numeric formula vocabulary.
        # This proves reference/numeric wiring, NOT Excel lifecycle compatibility.
        import re
        s=duration_model();w=load_workbook(io.BytesIO(render(s,'Duration')));m=w['Main']
        plans=[r for r in range(3,m.max_row+1) if m.cell(r,5).value=='Plan']

        for r in plans:m.cell(r+1,12,1 if m.cell(r,3).value=='A1290' else 0)
        def numeric(expression):
            e=str(expression).lstrip('=').replace('$','')
            try:return float(e)
            except ValueError:pass
            if e.startswith('IFERROR('):e=e[8:-3]
            if '/' in e:
                left,right=e.split('/',1);return numeric(left)/numeric(right)
            if e.startswith('SUMPRODUCT('):
                a,b=e[11:-1].split(',');return sum(numeric(x.value or 0)*numeric(y.value or 0) for x,y in zip(cells(a),cells(b)))
            if e.startswith('SUM('):return sum(numeric(c.value or 0) for c in cells(e[4:-1]))
            if re.fullmatch('[A-Z]+[0-9]+',e):return numeric(m[e].value or 0)
            raise AssertionError('Unsupported test expression '+e)
        def cells(rng):return [c for row in m[rng] for c in row]
        actual=next(r for r in range(3,m.max_row+1) if m.cell(r,4).value=='Project Actual')
        plan=actual-1
        self.assertAlmostEqual(numeric(m.cell(actual,11).value),.25)
        self.assertAlmostEqual(numeric(m.cell(actual,12).value),.25)
        self.assertAlmostEqual(numeric(m.cell(plan,11).value),1)
        self.assertAlmostEqual(numeric(m['I1'].value),32)
        for r in range(3,m.max_row+1):
            if m.cell(r,5).value=='WBS':self.assertAlmostEqual(numeric(m.cell(r,11).value),1)
        self.assertAlmostEqual(sum(numeric(m.cell(r,10).value) for r in plans),1)

    def test_duration_format_regression(self):
        raw=(FIX/'msp_n8.xml').read_bytes()
        for fmt in ('3','5','7','9','11'):
            changed=raw.replace(b'<Summary>0</Summary>',f'<Summary>0</Summary><Duration>PT8H</Duration><DurationFormat>{fmt}</DurationFormat>'.encode())
            self.assertEqual(bases(normalize(parse_xml(changed),changed),'Duration'),[8,8])
        for fmt in ('4','6','8','10','12','19','21','39','999',''):
            changed=raw.replace(b'<Summary>0</Summary>',f'<Summary>0</Summary><Duration>PT8H</Duration><DurationFormat>{fmt}</DurationFormat>'.encode())
            with self.assertRaises(ValueError):bases(normalize(parse_xml(changed),changed),'Duration')
    def test_custom_id_requires_evidence(self):
        raw=(FIX/'msp_n8.xml').read_bytes()
        plain=normalize(parse_xml(raw),raw)
        self.assertEqual(plain.activities[0].source_id,'4')
        self.assertEqual(plain.activities[0].source_object_id,'4')
        mapped=normalize(parse_xml(raw),raw,msp_activity_id_field='188743731')
        self.assertEqual(mapped.activities[0].source_id,'A1290')
        self.assertEqual(mapped.activities[0].source_id_field,'ExtendedAttribute:188743731')
        declared=raw.replace(b'<Tasks>',b'<ExtendedAttributes><ExtendedAttribute><FieldID>188743731</FieldID><Alias>Activity ID</Alias></ExtendedAttribute></ExtendedAttributes><Tasks>')
        self.assertEqual(normalize(parse_xml(declared),declared).activities[0].source_id,'A1290')
        duplicate=raw.replace(b'A1310',b'A1290')
        with self.assertRaises(ValueError):normalize(parse_xml(duplicate),duplicate,msp_activity_id_field='188743731')

    def test_http_xml_to_download(self):
        server=HTTPServer(('127.0.0.1',0),Handler);t=threading.Thread(target=server.serve_forever);t.start()
        try:
            url='http://127.0.0.1:'+str(server.server_port)
            with urllib.request.urlopen(urllib.request.Request(url+'?method=Equal',data=(FIX/'msp_n8.xml').read_bytes())) as r:
                self.assertEqual(r.status,200);self.assertEqual(r.headers['Cache-Control'],'no-store');self.assertIn('Main',load_workbook(io.BytesIO(r.read())).sheetnames)
            with self.assertRaises(urllib.error.HTTPError):urllib.request.urlopen(urllib.request.Request(url,data=(FIX/'msp_n8.xml').read_bytes()))
        finally:server.shutdown();t.join();server.server_close()
if __name__=='__main__':unittest.main()
