"""F1 XLSX adapter. Excel coordinates/formulas are confined to this module.

PS 89a6b80 contracts: weekly incremental Main; month-of-cutoff aggregation;
last-nonblank monthly cumulative; DF-1 latest-prior display, cutoff mask;
MS-2 monthly percent-complete ranges. Reuses PS theme data, not desktop services.
"""
from collections import OrderedDict
from dataclasses import asdict
from datetime import date
from io import BytesIO
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter as col
from .export_theme import DEFAULT_ACTIVITY_DATA_PALETTE as AD, DEFAULT_TIMESCALE_PALETTE as TS

THEME = json.loads(Path(__file__).with_name('dashboard_theme.json').read_text())
C = THEME['colors']
FIRST = 14
HEADERS = ['Row Type', 'WBS', 'Activity Name', 'P/A', 'Activity ID', 'Outline Level',
           'Plan Start', 'Plan Finish', 'Progress Basis', 'Progress Weight',
           '% Complete', 'Internal ID', 'Source Duration (hours)']


def text(cell, value):
    """All imported text stays literal, including =, +, -, @ prefixes."""
    cell.value = str(value or '')
    cell.data_type = 's'


def total_formula(ref):
    return f'=IF(COUNT({ref})=0,"",SUM({ref}))'


def cumulative_formula(first, here, last):
    return f'=IF(OR(COUNT({first}:{here})=0,COUNT({here}:{last})=0),"",SUM({first}:{here}))'


def last_value_formula(ref):
    return f'=IF(COUNT({ref})=0,"",IFERROR(LOOKUP(2,1/({ref}<>""),{ref}),""))'


def setup_sheet(ws, title, periods):
    ws.sheet_view.showGridLines = False
    ws.merge_cells('A1:M1'); ws['A1'] = title
    ws['A1'].font = Font(name=THEME['font'], size=18, bold=True, color=C['navy'])
    ws.row_dimensions[1].height = 32
    for i, label in enumerate(HEADERS, 1):
        ws.cell(4, i, label)
    for i, dt in enumerate(periods, FIRST):
        ws.cell(4, i, dt).number_format = 'dd-mmm-yy'
        ws.cell(3, i, f'W{i-FIRST+1}')
        ws.column_dimensions[col(i)].width = 13
    for cell in ws[4]:
        cell.font = Font(name=THEME['font'], bold=True, color=C['white'], size=10)
        cell.fill = PatternFill('solid', fgColor=C['navy'])
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    widths = [16, 18, 40, 6, 18, 8, 13, 13, 15, 13, 13, 38, 17]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[col(i)].width = width
    ws.column_dimensions['L'].hidden = True
    ws.column_dimensions['F'].hidden = True
    ws.row_dimensions[4].height = 34
    ws.freeze_panes = 'N5'
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.protection.sheet = True
    ws.protection.formatRows = False
    ws.print_title_rows = '1:4'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def style_row(ws, row, kind, level, pa, last):
    palette = [AD.wbs_level_1_fill, AD.wbs_level_2_fill, AD.wbs_level_3_fill, AD.wbs_level_4_fill]
    fill = palette[min(max(level-1, 0), 3)] if kind == 'WBS' else C['light_gray'] if kind == 'Project Summary' else C['white']
    period_fill = (TS.activity_plan_fill if pa == 'P' else TS.activity_actual_fill)
    if kind == 'Project Summary':
        period_fill = TS.project_plan_fill if pa == 'P' else TS.project_actual_fill
    elif kind == 'WBS':
        period_fill = (TS.wbs_level_1_plan_fill if pa == 'P' else TS.wbs_level_1_actual_fill) if level == 1 else (TS.wbs_level_2_plan_fill if pa == 'P' else TS.wbs_level_2_actual_fill)
    for c in range(1, last+1):
        cell = ws.cell(row, c)
        cell.font = Font(name=THEME['font'], size=10, bold=kind != 'Activity',
                         color=C['white'] if c >= FIRST and kind != 'Activity' else C['text'])
        cell.fill = PatternFill('solid', fgColor=fill if c < FIRST else period_fill)
        cell.alignment = Alignment(vertical='center', wrap_text=c == 3)
        cell.border = Border(bottom=Side(style='hair', color=C['border']))
        if c in (10, 11) or c >= FIRST:
            cell.number_format = '0.00%'
        elif c in (7, 8):
            cell.number_format = 'dd-mmm-yy'
        elif c in (9, 13):
            cell.number_format = '#,##0.00'
    ws.cell(row, 3).alignment = Alignment(indent=min(level, 15), vertical='center', wrap_text=True)
    ws.row_dimensions[row].outlineLevel = min(level, 7)
    ws.row_dimensions[row].height = 26


def write_main(wb, schedule, settings, weights, periods, plans):
    ws = wb.active; ws.title = 'Main'
    project_name = dict(schedule.source_project_fields).get('Name', schedule.source_project)
    setup_sheet(ws, 'WEEKLY PROGRESS', periods)
    text(ws['C2'], project_name)
    ws['I2'] = f'{settings.method} basis (not money)'
    last = FIRST + len(periods)-1
    records, activity_rows = [], []
    children, tasks = {}, {}
    for w in schedule.wbs: children.setdefault(w.parent, []).append(w)
    for i, a in enumerate(schedule.activities): tasks.setdefault(a.wbs, []).append((i, a))

    def pair(kind, key, path, name, level, activity=None, index=None):
        r = ws.max_row + 1
        for rr, pa in ((r, 'P'), (r+1, 'A')):
            values = [kind, path, name, pa, activity.source_id if activity else '', level,
                      activity.start if activity else None, activity.finish if activity else None,
                      weights[index] if activity else None, None, None, key,
                      activity.duration_hours if activity else None]
            for c, v in enumerate(values, 1):
                if isinstance(v, str): text(ws.cell(rr, c), v)
                else: ws.cell(rr, c, v)
            style_row(ws, rr, kind, level, pa, last)
            ws.cell(rr, 11, total_formula(f'{col(FIRST)}{rr}:{col(last)}{rr}'))
        if activity:
            for j, v in enumerate(plans[index], FIRST):
                ws.cell(r, j, v)
                ws.cell(r+1, j).protection = Protection(locked=False)
            activity_rows.append((r, index))
        records.append((r, kind, level, path, key))
        return r

    pair('Project Summary', schedule.import_id, '', project_name, 0)
    summaries = [(5, None)]

    def walk(parent=None, level=1, path=''):
        for w in children.get(parent, []):
            code = w.code if schedule.source_system == 'MSP' else '.'.join(filter(None, (path, w.code)))
            r = pair('WBS', w.key, code, w.name, level)
            walk(w.key, level+1, code)
            summaries.append((r, ws.max_row))
        for i, a in tasks.get(parent, []):
            pair('Activity', a.key, path, a.name, level, a, i)
    walk()
    end = ws.max_row
    for r, stop in summaries:
        lo, hi = r+2, stop or end
        for rr, pa in ((r, 'P'), (r+1, 'A')):
            if hi < lo:
                ws.cell(rr, 9, 0)
                continue
            types, pas = f'$A${lo}:$A${hi}', f'$D${lo}:$D${hi}'
            bases = f'$I${lo}:$I${hi}'
            ws.cell(rr, 9, f'=SUMIFS({bases},{types},"Activity",{pas},"{pa}")')
            for c in range(FIRST, last+1):
                values = f'{col(c)}{lo}:{col(c)}{hi}'
                ws.cell(rr, c, f'=IF(COUNTIFS({types},"Activity",{pas},"{pa}",{values},"<>")=0,"",'
                              f'IFERROR(SUMPRODUCT(--({types}="Activity"),--({pas}="{pa}"),{bases},{values})/$I{rr},""))')
    for r, _, _, _, _ in records:
        for rr in (r, r+1): ws.cell(rr, 10, f'=I{rr}/$I$5')
    actual_validation = DataValidation(type='decimal', operator='between', formula1=0, formula2=1, allow_blank=True)
    actual_validation.showErrorMessage = True
    actual_validation.errorTitle = 'Weekly increment'
    actual_validation.error = 'Enter a fraction from 0 to 1 (for example 5%), or leave blank.'
    actual_validation.showInputMessage = True
    actual_validation.prompt = 'Increment this week, not cumulative progress. Blank means not reported.'
    ws.add_data_validation(actual_validation)
    for r, _ in activity_rows:
        actual_validation.add(f'{col(FIRST)}{r+1}:{col(last)}{r+1}')
    ws.conditional_formatting.add(f'K5:K{end}', FormulaRule(formula=['K5>1'],
        fill=PatternFill('solid', fgColor=C['light_red']), font=Font(color=C['red'])))
    scurve = end+3
    for offset, label in enumerate(('Period Plan', 'Cumulative Plan', 'Period Actual', 'Cumulative Actual')):
        ws.cell(scurve+offset, 3, label)
        for c in range(FIRST, last+1):
            letter = col(c)
            if offset in (0, 2):
                src = f'{letter}{5 if offset == 0 else 6}'
                formula = f'=IF({src}="","",{src})'
            else:
                rr = scurve+offset-1
                formula = cumulative_formula(f'{col(FIRST)}{rr}', f'{letter}{rr}', f'{col(last)}{rr}')
            ws.cell(scurve+offset, c, formula).number_format = '0.00%'
            ws.cell(scurve+offset, c).fill = PatternFill('solid', fgColor=TS.scurve_acc_fill)
    return ws, records, activity_rows, scurve


def write_monthly(wb, main, records, periods, scurve):
    months = OrderedDict()
    for i, d in enumerate(periods): months.setdefault((d.year, d.month), []).append(i)
    ws = wb.create_sheet('Monthly')
    setup_sheet(ws, 'MONTHLY PROGRESS', [periods[v[-1]] for v in months.values()])
    ws['C2'] = '=Main!C2'; ws['I2'] = '=Main!I2'
    for i in range(len(months)): ws.cell(3, FIRST+i, f'M{i+1}')
    for r, kind, level, _, _ in records:
        for rr in (r, r+1):
            for c in range(1, FIRST):
                source = main.cell(rr, c)
                if c == 11: continue
                if source.value is not None: ws.cell(rr, c, f'=Main!{col(c)}{rr}')
            style_row(ws, rr, kind, level, 'P' if rr == r else 'A', FIRST+len(months)-1)
            ws.cell(rr, 11, total_formula(f'{col(FIRST)}{rr}:{col(FIRST+len(months)-1)}{rr}'))
            for j, indices in enumerate(months.values(), FIRST):
                lo, hi = FIRST+indices[0], FIRST+indices[-1]
                ref = f"'Main'!{col(lo)}{rr}:{col(hi)}{rr}"
                if kind == 'Activity' and rr == r:
                    values = [main.cell(rr, c).value for c in range(lo, hi+1)]
                    nums = [v for v in values if isinstance(v, (int, float))]
                    if nums: ws.cell(rr, j, sum(nums))
                else:
                    ws.cell(rr, j, total_formula(ref))
    for offset in range(4):
        rr = scurve+offset
        ws.cell(rr, 3, main.cell(rr, 3).value)
        for j, indices in enumerate(months.values(), FIRST):
            ref = f"'Main'!{col(FIRST+indices[0])}{rr}:{col(FIRST+indices[-1])}{rr}"
            ws.cell(rr, j, last_value_formula(ref) if offset in (1, 3) else total_formula(ref)).number_format = '0.00%'
    ws.conditional_formatting.add(f'K5:K{scurve-3}', FormulaRule(formula=['K5>1'], fill=PatternFill('solid', fgColor=C['light_red'])))
    return ws, list(months.values())


def write_amount(wb, records, main):
    ws = wb.create_sheet('Activity Amount')
    ws.append(['ACTIVITY AMOUNT — allocated Contract Value'])
    ws.append(['Input only. Blank is effective zero; entering money never changes Equal/Duration.'])
    ws.append(['WBS', 'Activity ID', 'Activity Name', 'Amount', 'Internal ID', 'Effective Amount', 'Input state'])
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1=0, allow_blank=True)
    dv.showErrorMessage = True; dv.error = 'Enter nonnegative allocated Contract Value or leave blank.'
    ws.add_data_validation(dv)
    for r, kind, level, _, key in records:
        if kind == 'Project Summary': continue
        ar = ws.max_row+1
        ws.append([f'=Main!B{r}', f'=IF(Main!E{r}="","",Main!E{r})', f'=Main!C{r}', None, key])
        text(ws.cell(ar, 5), key)
        ws.row_dimensions[ar].outlineLevel = min(level, 7)
        ws.row_dimensions[ar].height = 25
        if kind == 'Activity':
            ws.cell(ar, 4).protection = Protection(locked=False)
            ws.cell(ar, 4).fill = PatternFill('solid', fgColor=TS.activity_actual_fill)
            ws.cell(ar, 4).number_format = '#,##0.00'
            ws.cell(ar, 6, f'=IF(D{ar}="",0,D{ar})')
            ws.cell(ar, 7, f'=IF(D{ar}="","Missing",IF(D{ar}=0,"Zero","Populated"))')
            dv.add(ws.cell(ar, 4))
            ws.conditional_formatting.add(f'D{ar}', FormulaRule(formula=[f'OR(D{ar}="",D{ar}=0)'], fill=PatternFill('solid', fgColor=C['light_amber'])))
        else:
            for cell in ws[ar]:
                cell.fill = PatternFill('solid', fgColor=AD.wbs_level_1_fill)
                cell.font = Font(bold=True, name=THEME['font'])
    for cell in ws[3]:
        cell.fill = PatternFill('solid', fgColor=C['navy']); cell.font = Font(color=C['white'], bold=True)
    for key, width in [('A', 22), ('B', 22), ('C', 50), ('D', 20), ('F', 20), ('G', 16)]:
        ws.column_dimensions[key].width = width
    ws.column_dimensions['E'].hidden = True
    ws.freeze_panes = 'D4'; ws.sheet_view.showGridLines = False
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.protection.sheet = True; ws.protection.formatRows = False


def chart(data, date_col, plan_col, actual_col, marker_plan, marker_actual, line_col, last, title):
    ch = LineChart(); ch.title = title
    ch.style = THEME['layout']['chart_style']
    ch.width = THEME['layout']['chart_width']; ch.height = THEME['layout']['chart_height']
    ch.y_axis.numFmt = '0%'; ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1.1
    ch.y_axis.majorUnit = .25; ch.y_axis.title = 'Cumulative progress'
    # PS _date_axis_for_line_chart: preserve LineChart's reciprocal axis IDs.
    # DateAxis defaults to 500, leaving y.crossAx=10 dangling. Correcting this
    # reference alone does not establish Desktop Excel compatibility.
    ch.x_axis = DateAxis(axId=10, crossAx=100)
    ch.y_axis.axId = 100; ch.y_axis.crossAx = 10
    ch.x_axis.number_format = 'mmm-yy'
    ch.x_axis.title = 'Reporting date'; ch.display_blanks = 'gap'
    ch.visible_cells_only = False
    for index, c in enumerate((plan_col, actual_col, marker_plan, marker_actual, line_col)):
        ch.add_data(Reference(data, min_col=c, min_row=1, max_row=last), titles_from_data=True)
        s = ch.series[-1]
        color = C['blue'] if index in (0, 2) else C['green'] if index in (1, 3) else C['red']
        s.graphicalProperties.line.width = 26000
        # DrawingML line fills are mutually exclusive. Assign solidFill only
        # to curve lines; setting noFill later does not clear it in openpyxl.
        if index in (2, 3):
            s.graphicalProperties.line.noFill = True
            s.marker.symbol = 'circle'; s.marker.size = 6
            s.marker.graphicalProperties.solidFill = color
            s.marker.graphicalProperties.line.solidFill = color
        elif index == 4:
            s.graphicalProperties.line.noFill = True
            gp = GraphicalProperties(); gp.line.solidFill = C['red']; gp.line.prstDash = 'dash'
            s.marker.symbol = 'none'
            s.errBars = ErrorBars(errDir='y', errBarType='minus', errValType='fixedVal', val=1, noEndCap=True, spPr=gp)
            # PS cutoff label uses the live category, not a cached date string.
            s.dLbls = DataLabelList(showVal=False, showCatName=True, showSerName=False,
                                   showLegendKey=False, dLblPos='t', numFmt='"Cutoff "dd/mm/yyyy')
        else:
            s.graphicalProperties.line.solidFill = color
            s.marker.symbol = 'none'
    ch.set_categories(Reference(data, min_col=date_col, min_row=2, max_row=last))
    from openpyxl.chart.legend import LegendEntry
    ch.legend.legendEntry = [LegendEntry(idx=i, delete=True) for i in (2, 3, 4)]
    ch.legend.position = 'b'
    return ch


def write_dashboard(wb, main, monthly, periods, months, scurve):
    progress = wb.create_sheet('progress')
    progress.append(['Date', 'Plan', 'Actual'])
    for i, d in enumerate(periods, 2):
        c = col(FIRST+i-2)
        progress.append([d, f'=Main!{c}{scurve+1}', f'=IF(Main!{c}{scurve+3}="","",Main!{c}{scurve+3})'])
    progress.sheet_state = 'hidden'
    data = wb.create_sheet('Dashboard_Data')
    data.append(['Weekly Date', 'Weekly Plan', 'Weekly Actual', 'Monthly Date', 'Monthly Plan', 'Monthly Actual',
                 'Selected Date', 'Selected Plan', 'Selected Actual', 'Weekly Cutoff', 'Monthly Cutoff',
                 'Selected Actual Raw', 'Marker Date', 'Plan marker', 'Actual marker'])
    dash = wb.create_sheet('Dashboard', 0)
    dash.sheet_view.showGridLines = False
    dash.merge_cells('B2:M3'); dash['B2'] = 'PROGRESS DASHBOARD'
    dash['B2'].font = Font(name=THEME['font'], size=24, bold=True, color=C['navy'])
    dash['B4'] = '=Main!C2'
    dash['F5'] = 'View'; dash['G5'] = THEME['layout']['default_view']
    dash['J5'] = 'Actual cutoff'; dash['K5'] = periods[-1]; dash['K5'].number_format = 'dd-mmm-yyyy'
    dash.merge_cells('B7:M7'); dash['B7'] = 'Weekly Actual is entered in Main. F9 recalculates. Cutoff changes display only.'
    dash['B7'].font = Font(size=10, color=C['muted'])
    last = len(periods)+1
    for i, d in enumerate(periods, 2):
        data.cell(i, 1, f'=progress!A{i}')
        data.cell(i, 2, f'=progress!B{i}')
        data.cell(i, 3, f'=IF(progress!C{i}="","",progress!C{i})')
        data.cell(i, 10, d)
        if i-2 < len(months):
            src = months[i-2][-1]+2
            data.cell(i, 4, periods[src-2])
            data.cell(i, 5, f'=progress!B{src}')
            data.cell(i, 6, f'=IF(progress!C{src}="","",progress!C{src})')
            data.cell(i, 11, periods[src-2])
        data.cell(i, 7, f'=IF(Dashboard!$G$5="Weekly",IF(A{i}="","",A{i}),IF(D{i}="","",D{i}))')
        data.cell(i, 8, f'=IF(G{i}="",NA(),IF(Dashboard!$G$5="Weekly",B{i},E{i}))')
        # Exact DF-1 separation: raw display lookup is not cutoff-dependent;
        # zero fallback is display-only, marker uses actual-presence separately.
        data.cell(i, 12, f'=IF(G{i}="","",IFERROR(LOOKUP(2,1/(($J$2:$J${last}<=G{i})*($C$2:$C${last}<>"")),$C$2:$C${last}),0))')
        data.cell(i, 9, f'=IF(G{i}="",NA(),IF(G{i}>Dashboard!$K$5,NA(),IF(L{i}="",NA(),L{i})))')
        data.cell(i, 13, f'=IF(G{i}="","",G{i})')
        selected = f'IFERROR(LOOKUP(2,1/(($G$2:$G${last}<>"")*($G$2:$G${last}<=Dashboard!$K$5)),$G$2:$G${last}),0)'
        data.cell(i, 14, f'=IF(OR(G{i}="",G{i}<>{selected}),NA(),H{i})')
        data.cell(i, 15, f'=IFERROR(IF(OR(G{i}="",G{i}<>{selected}),NA(),IF(AND(COUNTIFS($J$2:$J${last},"<="&G{i},$C$2:$C${last},"<>")>0,ABS(L{i}-H{i})>0.0000001),L{i},NA())),NA())')
        data.cell(i, 28, f'=IF(OR(G{i}="",G{i}<>{selected}),NA(),1)')
        for c in (1, 4, 7, 10, 11, 13): data.cell(i, c).number_format = 'dd-mmm-yyyy'
    data.cell(1, 28, 'Dashboard cutoff indicator (exclusive)')
    # KPI resolves selected category <= cutoff, the same category as chart marker.
    dash['B10'] = f'=IFERROR(LOOKUP(2,1/((Dashboard_Data!$G$2:$G${last}<>"")*(Dashboard_Data!$G$2:$G${last}<=K5)),Dashboard_Data!$H$2:$H${last}),0)'
    dash['F10'] = f'=IFERROR(LOOKUP(2,1/((Dashboard_Data!$G$2:$G${last}<>"")*(Dashboard_Data!$G$2:$G${last}<=K5)),Dashboard_Data!$L$2:$L${last}),0)'
    dash['J10'] = '=F10-B10'
    for label, title, val in [('B9', 'PLANNED', 'B10'), ('F9', 'ACTUAL', 'F10'), ('J9', 'PROGRESS GAP', 'J10')]:
        dash[label] = title; dash[label].font = Font(name=THEME['font'], bold=True, color=C['muted'])
        dash[val].number_format = '0.00%'; dash[val].font = Font(name=THEME['font'], size=26, bold=True, color=C['navy'])
    dv = DataValidation(type='list', formula1='"Weekly,Monthly"'); dv.showErrorMessage = True
    dash.add_data_validation(dv); dv.add(dash['G5'])
    cd = DataValidation(type='list', formula1=f'INDIRECT(IF($G$5="Weekly","Dashboard_Data!$J$2:$J${last}","Dashboard_Data!$K$2:$K${len(months)+1}"))')
    cd.showErrorMessage = False  # PS indicator accepts dates between reporting points too.
    dash.add_data_validation(cd); cd.add(dash['K5'])
    for cell in ('G5', 'K5'):
        dash[cell].protection = Protection(locked=False)
        dash[cell].fill = PatternFill('solid', fgColor=C['light_blue'])
    for c in range(1, 15): dash.column_dimensions[col(c)].width = 12
    dash.row_dimensions[10].height = 40
    dash.add_chart(chart(data, 7, 8, 9, 14, 15, 28, last, 'Plan and Actual'), 'B13')
    dash['B33'] = 'Weights are not money. Activity Amount does not change this dashboard.'
    dash.protection.sheet = True
    data.sheet_state = 'hidden'
    # Independent Main/Monthly overlays reuse the same raw weekly authority.
    for ws, points, offset in ((main, list(range(len(periods))), 30), (monthly, [m[-1] for m in months], 38)):
        ws['B2'] = 'Actual cutoff'; ws['G2'] = periods[-1]; ws['G2'].number_format = 'dd-mmm-yyyy'
        ws['G2'].protection = Protection(locked=False)
        selector = DataValidation(type='list', formula1=f'INDIRECT("Dashboard_Data!${"J" if ws is main else "K"}$2:${"J" if ws is main else "K"}${len(points)+1}")')
        ws.add_data_validation(selector); selector.add(ws['G2'])
        names = ['Date', 'Plan', 'Actual', 'Plan marker', 'Actual marker', 'Cutoff']
        for k, n in enumerate(names): data.cell(1, offset+k, ws.title+' '+n)
        dc, pc, ac, mp, ma, lc = [col(offset+k) for k in range(6)]
        for n, idx in enumerate(points, 2):
            data.cell(n, offset, periods[idx])
            data.cell(n, offset+1, f'=progress!B{idx+2}')
            lookup = f'IFERROR(LOOKUP(2,1/(($J$2:$J${last}<={dc}{n})*($C$2:$C${last}<>"")),$C$2:$C${last}),0)'
            data.cell(n, offset+2, f'=IF({dc}{n}>{ws.title}!$G$2,NA(),{lookup})')
            anchor = f'IFERROR(LOOKUP(2,1/(${dc}$2:${dc}${len(points)+1}<={ws.title}!$G$2),${dc}$2:${dc}${len(points)+1}),0)'
            data.cell(n, offset+3, f'=IF({dc}{n}={anchor},{pc}{n},NA())')
            data.cell(n, offset+4, f'=IF(AND({dc}{n}={anchor},COUNTIFS($J$2:$J${last},"<="&{dc}{n},$C$2:$C${last},"<>")>0),{ac}{n},NA())')
            data.cell(n, offset+5, f'=IF({dc}{n}={anchor},1,NA())')
        ws.add_chart(chart(data, offset, offset+1, offset+2, offset+3, offset+4, offset+5,
                           len(points)+1, ws.title+' S-curve'), f'B{scurve+6}')


def render(schedule, settings, prepared):
    weights, periods, plans, methods, warnings = prepared
    wb = Workbook()
    main, records, activities, scurve = write_main(wb, schedule, settings, weights, periods, plans)
    monthly, months = write_monthly(wb, main, records, periods, scurve)
    write_amount(wb, records, main)
    write_dashboard(wb, main, monthly, periods, months, scurve)
    guide = wb.create_sheet('Guide')
    guide.append(['BLUEBIRD — PROGRESS WORKBOOK'])
    for message in [
        'Enter weekly incremental Actual in unlocked Main cells. Use 5% or 0.05, not cumulative values.',
        'Blank means not reported; zero is a recorded zero. Correct earlier cells directly if necessary.',
        'F9 recalculates formulas. Save, close and reopen in Desktop Excel; no server rebuild is needed for Actual.',
        'Manual calculation; calculate-on-save and first-open recalculation requested. Excel settings can override this.',
        'Monthly assigns each week to the month of its cutoff date; it does not split across month boundaries.',
        'Dashboard carries the latest Actual for display up to its cutoff, without filling blank source records.',
        'Activity Amount is allocated Contract Value, not cost/BAC. It is stored separately and does not change weights.',
        'Milestones have zero Progress Weight; their entered Contract Value is retained.',
        'Amount weighting and workbook refresh are later capabilities. Do not restructure rows or edit calculated Plan.',
        f'Creation: {settings.method}; cutoff {settings.cutoff}; distribution {settings.distribution}.',
    ] + warnings:
        guide.append([message])
    guide.column_dimensions['A'].width = 110
    for row in guide:
        row[0].alignment = Alignment(wrap_text=True, vertical='center')
        row[0].font = Font(name=THEME['font'], size=11)
        guide.row_dimensions[row[0].row].height = 36
    meta = wb.create_sheet('_Metadata')
    for k, v in [('schema', 'bluebird-f1-1'), ('weight_method', settings.method),
                 ('weekly_cutoff', settings.cutoff), ('plan_distribution', settings.distribution),
                 ('ps_reference', '89a6b80'), ('import_id', schedule.import_id),
                 ('source_hash', schedule.source_hash), ('source_system', schedule.source_system),
                 ('source_project', schedule.source_project), ('amount_basis', 'allocated Contract Value')]:
        meta.append([k, v])
    for i, a in enumerate(schedule.activities):
        meta.append(['activity', json.dumps({**asdict(a), 'distribution': methods[i]}, default=str, ensure_ascii=False)])
    for w in schedule.wbs: meta.append(['wbs', json.dumps(asdict(w), ensure_ascii=False)])
    for row in meta:
        for cell in row: cell.data_type = 's'
    meta.sheet_state = 'veryHidden'
    wb.calculation = CalcProperties(calcMode='manual', calcOnSave=True,
                                    fullCalcOnLoad=True, forceFullCalc=True, calcId=0)
    out = BytesIO()
    wb.save(out); wb.close()
    return out.getvalue()
