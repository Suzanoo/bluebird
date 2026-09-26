"""Serialized chart semantics: XML well-formedness alone cannot detect crossAx bugs."""
from dataclasses import replace
from io import BytesIO
import posixpath
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from engine.domain import Settings, prepare
from engine.service import convert
from engine.workbook import render
from test_f1 import schedule
from test_runtime import fixture

C = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
XDR = 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'


def check_line_fills(case, root):
    # CT_LineProperties/EG_LineFillProperties is a choice, not a sequence.
    # openpyxl permits both assignments and serializes both without validation.
    choices = {f'{{{A}}}{name}' for name in ('noFill', 'solidFill', 'gradFill', 'pattFill')}
    for line in root.iter(f'{{{A}}}ln'):
        case.assertLessEqual(sum(child.tag in choices for child in line), 1,
                             'mutually exclusive DrawingML line fills')


def check_axes(case, root):
    # Read definitions, not every axId (lineChart contains references as well).
    plot = root.find(f'.//{{{C}}}plotArea')
    axes = [node for node in plot if node.tag in
            {f'{{{C}}}{kind}' for kind in ('dateAx', 'valAx', 'catAx', 'serAx')}]
    ids = [node.find(f'{{{C}}}axId').get('val') for node in axes]
    case.assertEqual(len(ids), len(set(ids)), 'duplicate axis definitions')
    crossings = {node.find(f'{{{C}}}axId').get('val'):
                 node.find(f'{{{C}}}crossAx').get('val') for node in axes}
    for own, other in crossings.items():
        case.assertIn(other, ids, 'dangling crossing axis')
        case.assertNotEqual(own, other)
        case.assertEqual(crossings[other], own, 'non-reciprocal crossing axes')
    case.assertEqual(set(ids), {'10', '100'})
    line = plot.find(f'{{{C}}}lineChart')
    case.assertEqual({n.get('val') for n in line.findall(f'{{{C}}}axId')}, set(ids))
    case.assertEqual(len(line.findall(f'{{{C}}}ser')), 5, 'all F1 series retained')


def check_package(case, raw):
    with ZipFile(BytesIO(raw)) as package:
        case.assertIsNone(package.testzip())
        names = set(package.namelist())
        types = {n.get('PartName').lstrip('/'): n.get('ContentType')
                 for n in ET.fromstring(package.read('[Content_Types].xml'))
                 if n.tag == f'{{{CT}}}Override'}
        charts = sorted(n for n in names if n.startswith('xl/charts/chart') and n.endswith('.xml'))
        drawings = sorted(n for n in names if n.startswith('xl/drawings/drawing') and n.endswith('.xml'))
        case.assertEqual(len(charts), 3)
        case.assertEqual(len(drawings), 3)
        for name in charts:
            case.assertEqual(types[name], 'application/vnd.openxmlformats-officedocument.drawingml.chart+xml')
            check_axes(case, ET.fromstring(package.read(name)))
        # Resolve ownership through package relationships, not drawing numbering.
        def target(base, value):
            return value.lstrip('/') if value.startswith('/') else posixpath.normpath(posixpath.join(base, value))
        bookrels = {n.get('Id'): target('xl', n.get('Target')) for n in
                    ET.fromstring(package.read('xl/_rels/workbook.xml.rels'))}
        drawing_sheets = {}
        for sheet in ET.fromstring(package.read('xl/workbook.xml')).findall('.//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet'):
            part = bookrels[sheet.get(f'{{{R}}}id')]
            relpath = posixpath.dirname(part) + '/_rels/' + posixpath.basename(part) + '.rels'
            if relpath in names:
                for rel in ET.fromstring(package.read(relpath)):
                    if rel.get('Type').endswith('/drawing'):
                        drawing_sheets[target(posixpath.dirname(part), rel.get('Target'))] = sheet.get('name')
        for name in drawings:
            case.assertEqual(types[name], 'application/vnd.openxmlformats-officedocument.drawing+xml')
            root = ET.fromstring(package.read(name))
            case.assertEqual(len(root), 1)
            anchor = root[0]
            expected = 'oneCellAnchor' if drawing_sheets[name] == 'Dashboard' else 'twoCellAnchor'
            case.assertEqual(anchor.tag, f'{{{XDR}}}{expected}')
            for coordinate in ('col', 'row', 'colOff', 'rowOff'):
                case.assertGreaterEqual(int(anchor.find(f'{{{XDR}}}from/{{{XDR}}}{coordinate}').text), 0)
            if expected == 'oneCellAnchor':
                ext = anchor.find(f'{{{XDR}}}ext')
                case.assertGreater(int(ext.get('cx')), 0)
                case.assertGreater(int(ext.get('cy')), 0)
            else:
                case.assertEqual(anchor.get('editAs'), 'twoCell')
                for coordinate in ('col', 'row'):
                    case.assertGreater(int(anchor.find(f'{{{XDR}}}to/{{{XDR}}}{coordinate}').text),
                                       int(anchor.find(f'{{{XDR}}}from/{{{XDR}}}{coordinate}').text))
            frame = anchor.find(f'{{{XDR}}}graphicFrame')
            case.assertIsNotNone(frame.find(f'{{{XDR}}}xfrm'))
            ref = frame.find(f'.//{{{C}}}chart').get(f'{{{R}}}id')
            relpath = posixpath.dirname(name) + '/_rels/' + posixpath.basename(name) + '.rels'
            rels = {n.get('Id'): n for n in ET.fromstring(package.read(relpath))}
            rel = rels[ref]
            case.assertTrue(rel.get('Type').endswith('/chart'))
            target = rel.get('Target')
            resolved = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(name), target))
            case.assertIn(resolved, charts)


class ChartIntegrityTests(unittest.TestCase):
    def test_both_sources_and_methods_have_valid_serialized_charts(self):
        for source in ('msp', 'p6'):
            for method in ('Equal', 'Duration'):
                with self.subTest(source=source, method=method):
                    raw, _ = convert(fixture(source, True), method)
                    check_package(self, raw)

    def test_large_row_anchor_and_openpyxl_roundtrip(self):
        original = schedule()
        s = replace(original, activities=tuple(
            replace(original.activities[0], key=f'a{i}', source_id=f'A{i}')
            for i in range(300)))
        settings = Settings('Duration')
        raw = render(s, settings, prepare(s, settings))
        check_package(self, raw)
        wb = load_workbook(BytesIO(raw))
        out = BytesIO()
        wb.save(out)
        wb.close()
        check_package(self, out.getvalue())

    def test_historical_dangling_cross_axis_is_rejected(self):
        raw, _ = convert(fixture('msp', True), 'Equal')
        with ZipFile(BytesIO(raw)) as package:
            root = ET.fromstring(package.read('xl/charts/chart1.xml'))
        check_axes(self, root)
        # The exact uploaded failure: date axis 500, value crossAx still 10.
        root.find(f'.//{{{C}}}dateAx/{{{C}}}axId').set('val', '500')
        for node in root.findall(f'.//{{{C}}}lineChart/{{{C}}}axId'):
            if node.get('val') == '10':
                node.set('val', '500')
        with self.assertRaises(AssertionError):
            check_axes(self, root)

    def test_line_fill_choice_and_all_series_survive_save_roundtrip(self):
        for source in ('msp', 'p6'):
            for method in ('Equal', 'Duration'):
                with self.subTest(source=source, method=method):
                    raw, _ = convert(fixture(source, True), method)
                    for iteration in range(2):
                        with ZipFile(BytesIO(raw)) as package:
                            for number in (1, 2, 3):
                                root = ET.fromstring(package.read(f'xl/charts/chart{number}.xml'))
                                check_line_fills(self, root)
                                check_axes(self, root)
                                series = root.findall(f'.//{{{C}}}lineChart/{{{C}}}ser')
                                for index, s in enumerate(series):
                                    line = s.find(f'{{{C}}}spPr/{{{A}}}ln')
                                    expected = 'solidFill' if index < 2 else 'noFill'
                                    self.assertIsNotNone(line.find(f'{{{A}}}{expected}'))
                                    self.assertIsNotNone(s.find(f'{{{C}}}cat'))
                                    self.assertIsNotNone(s.find(f'{{{C}}}val/{{{C}}}numRef/{{{C}}}f'))
                                for index, color in ((2, '2F75B5'), (3, '70AD47')):
                                    marker = series[index].find(f'{{{C}}}marker')
                                    self.assertEqual(marker.find(f'{{{C}}}symbol').get('val'), 'circle')
                                    self.assertEqual(marker.find(f'{{{C}}}spPr/{{{A}}}solidFill/{{{A}}}srgbClr').get('val'), color)
                                error_line = series[4].find(f'{{{C}}}errBars/{{{C}}}spPr/{{{A}}}ln')
                                self.assertEqual(error_line.find(f'{{{A}}}solidFill/{{{A}}}srgbClr').get('val'), 'C00000')
                                self.assertEqual(error_line.find(f'{{{A}}}prstDash').get('val'), 'dash')
                                self.assertEqual(series[4].find(f'{{{C}}}dLbls/{{{C}}}showCatName').get('val'), '1')
                        if iteration == 0:
                            wb = load_workbook(BytesIO(raw))
                            out = BytesIO()
                            wb.save(out); wb.close()
                            raw = out.getvalue()

    def test_historical_conflicting_line_fills_are_rejected(self):
        # Tiny literal reproducer independent of renderer assignments.
        bad = ET.fromstring(f'<a:ln xmlns:a="{A}"><a:noFill/>'
                            '<a:solidFill><a:srgbClr val="2F75B5"/></a:solidFill></a:ln>')
        with self.assertRaisesRegex(AssertionError, 'mutually exclusive'):
            check_line_fills(self, bad)
        bad.remove(bad.find(f'{{{A}}}solidFill'))
        check_line_fills(self, bad)


if __name__ == '__main__':
    unittest.main()
