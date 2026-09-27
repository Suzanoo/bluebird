"""Trial #1: source-order and display-only regressions; not Excel acceptance."""
from dataclasses import replace
from io import BytesIO
from uuid import UUID, uuid5
from xml.etree import ElementTree as ET
import json
import unittest
from openpyxl import load_workbook
from model import normalize, parse_xml, ordered_children, Wbs, sequence_number
from engine.domain import Settings, prepare
from engine.amount import preview_amounts
from engine.workbook import render, FIRST, total_formula
from test_f2 import amount_xml, selected
from test_f1_chart_integrity import check_package, check_line_fills
from zipfile import ZipFile

IID = '12345678-1234-5678-1234-567812345678'


def structure_xml(source):
    root = parse_xml(amount_xml(source))
    p = root if source == 'msp' else root.find('Project')
    def add(parent, tag, **fields):
        node=ET.SubElement(parent,tag)
        for k,v in fields.items(): ET.SubElement(node,k).text=str(v)
        return node
    if source == 'p6':
        p.find('WBS/Code').text='Z-first'
        add(p.find('WBS'),'SequenceNumber').text='20'
        # Encounter order Z, A, nested; numeric/alphabetic code sort is incorrect.
        add(p,'WBS',ObjectId='20',Code='99-first',Name='First root',SequenceNumber='10')
        add(p,'WBS',ObjectId='30',Code='child',Name='Nested',ParentObjectId='20',SequenceNumber='1')
        for a,w in zip(p.findall('Activity'),('20','30','20','10')):a.find('WBSObjectId').text=w
    else:
        tasks=p.find('Tasks'); original=list(tasks);tasks.clear()
        summary=original[0];summary.find('WBS').text='Z-first';tasks.append(summary)
        tasks.append(original[1])
        add(tasks,'Task',UID='20',ID='20',Name='Nested',Summary='1',OutlineLevel='2',WBS='Z-first.child')
        original[2].find('OutlineLevel').text='3';tasks.append(original[2]);tasks.append(original[3])
        add(tasks,'Task',UID='30',ID='30',Name='Last root',Summary='1',OutlineLevel='1',WBS='A-last')
        tasks.append(original[4])
    return ET.tostring(root)


def normalized(source):
    raw=structure_xml(source)
    return normalize(parse_xml(raw),raw,import_id=IID,production=True)


def make(source,method='Amount'):
    s=normalized(source);settings=Settings(method,distribution='flat',amount_field=selected(source) if method=='Amount' else None)
    selection=preview_amounts(s,selected(source)).require_valid() if method=='Amount' else None
    return s,render(s,settings,prepare(s,settings,selection),selection)


def rows(ws):
    return [(r,ws.cell(r,12).value,ws.cell(r,1).value) for r in range(7,ws.max_row+1) if ws.cell(r,4).value=='P']


class StructureTests(unittest.TestCase):
    def test_p6_sequence_nested_direct_activity_placement(self):
        s,raw=make('p6');w=load_workbook(BytesIO(raw));ws=w['Main']
        self.assertEqual([ws.cell(r,3).value for r,_,_ in rows(ws)],
            ['First root','Work 1','Work 3','Nested','Work 2','Structure','Work 4'])
        self.assertEqual(ws['B11'].value,'99-first')
        self.assertEqual(ws['B15'].value,'99-first.child')
        self.assertIn('$A$9:$A$16',ws['I7'].value)
        self.assertIn('$A$15:$A$16',ws['I13'].value)
        w.close()

    def test_msp_shared_encounter_order_and_outline(self):
        s,raw=make('msp');w=load_workbook(BytesIO(raw));ws=w['Main']
        self.assertEqual([ws.cell(r,3).value for r,_,_ in rows(ws)],
            ['Structure','Work 1','Nested','Work 2','Work 3','Last root','Work 4'])
        self.assertEqual([ws.row_dimensions[r].outlineLevel for r,_,_ in rows(ws)],[1,2,2,3,2,1,2])
        self.assertIn('$A$9:$A$16',ws['I7'].value)
        self.assertIn('$A$13:$A$14',ws['I11'].value)
        w.close()

    def test_missing_tied_and_nonfinite_p6_sequence_fallback(self):
        nodes=[Wbs(str(i),str(i),c,c,None,i,n) for i,(c,n) in enumerate(
            [('Z',None),('99',2),('1',2),('A',None)])]
        self.assertEqual([x[2].code for x in ordered_children('P6',nodes,[])],['99','1','Z','A'])
        self.assertEqual([x[2].code for x in ordered_children('P6',[replace(x,sequence_number=None,source_order=None) for x in nodes],[])],['Z','99','1','A'])
        for v in ('','NaN','Infinity','bad'):self.assertIsNone(sequence_number(v))
        raw=structure_xml('p6').replace(b'<SequenceNumber>10</SequenceNumber>',b'<SequenceNumber>bad</SequenceNumber>')
        s=normalize(parse_xml(raw),raw,production=True)
        self.assertIsNone(next(x for x in s.wbs if x.source_id=='20').sequence_number)

    def test_legacy_msp_metadata_absent_uses_existing_groups_not_codes(self):
        s=normalized('msp'); roots=[replace(w,source_order=None) for w in s.wbs if w.parent is None]
        a=replace(s.activities[0],wbs=None,source_order=None)
        ordered=ordered_children('MSP',roots,[(0,a)])
        self.assertEqual([x[2].source_id for x in ordered],['10','30','1'])
        self.assertEqual(ordered[-1][1],0)

    def test_identity_independent_of_order_and_wbs_labels(self):
        for source in ('p6','msp'):
            s=normalized(source)
            for a in s.activities:self.assertEqual(a.key,str(uuid5(UUID(IID),'activity:object:'+a.source_object_id)))
            raw=structure_xml(source).replace(b'Z-first',b'RENAMED').replace(b'<SequenceNumber>10',b'<SequenceNumber>900')
            other=normalize(parse_xml(raw),raw,import_id=IID,production=True)
            self.assertEqual({a.source_object_id:a.key for a in s.activities},{a.source_object_id:a.key for a in other.activities})

    def test_views_weights_metadata_and_contract_value_stay_aligned(self):
        for source in ('p6','msp'):
            for method,expected in [('Equal',[1,1,1,0]),('Duration',[8,8,8,0]),('Amount',[100,300,0,0])]:
                s,raw=make(source,method);w=load_workbook(BytesIO(raw));main=w['Main'];monthly=w['Monthly'];amount=w['Activity Amount']
                rr=rows(main)
                self.assertEqual([amount.cell(i,5).value for i in range(4,amount.max_row+1)],[k for _,k,_ in rr])
                for i,(r,key,kind) in enumerate(rr,4):
                    self.assertEqual(monthly.cell(r,12).value,f'=Main!L{r}')
                    self.assertEqual(amount.cell(i,1).value,f'=Main!B{r}')
                    if kind=='Activity':
                        a=next(a for a in s.activities if a.key==key);ix=s.activities.index(a)
                        self.assertEqual(main.cell(r,9).value,expected[ix]);self.assertEqual(main.cell(r+1,9).value,expected[ix])
                        self.assertAlmostEqual(sum(main.cell(r,c).value or 0 for c in range(FIRST,FIRST+4)),1)
                        if method=='Amount':self.assertEqual(amount.cell(i,4).value,[100,300,0,500][ix])
                meta=[json.loads(row[1].value) for row in w['_Metadata'] if row[0].value in ('activity','wbs')]
                self.assertTrue(meta);self.assertTrue(all('source_order' in x for x in meta))
                w.close()

    def test_actual_display_only_and_roundtrip_blank_zero(self):
        for source in ('p6','msp'):
            s,raw=make(source);w=load_workbook(BytesIO(raw));main=w['Main']
            for sheet in ('Main','Monthly'):
                ws=w[sheet]
                for r,key,kind in rows(main):
                    for c in (1,2,3,5,7,8,13):
                        self.assertEqual(ws.cell(r+1,c).number_format,';;;')
                        self.assertNotEqual(ws.cell(r,c).number_format,';;;')
                    for c in (4,9,10,11,14):self.assertNotEqual(ws.cell(r+1,c).number_format,';;;')
            for r,key,kind in rows(main):
                self.assertEqual(main.cell(r,12).value,main.cell(r+1,12).value)
                if kind=='Activity':
                    self.assertEqual(main.cell(r,1).value,main.cell(r+1,1).value)
                    self.assertIsNone(main.cell(r+1,14).value);self.assertFalse(main.cell(r+1,14).protection.locked)
                    main.cell(r+1,15,0)
            out=BytesIO();w.save(out);w.close();w=load_workbook(BytesIO(out.getvalue()))
            for r,_,kind in rows(w['Main']):
                if kind=='Activity':
                    self.assertIsNone(w['Main'].cell(r+1,14).value);self.assertEqual(w['Main'].cell(r+1,15).value,0)
                    self.assertIn('COUNT',w['Main'].cell(r+1,11).value)
                    self.assertIn('COUNT',w['Monthly'].cell(r+1,14).value)
            self.assertFalse(any('Activity Amount' in str(c.value) for ws in (w['Main'],w['Monthly'],w['Dashboard']) for row in ws for c in row if c.data_type=='f'))
            w.close()

    def test_all_sources_methods_keep_three_valid_charts(self):
        for source in ('p6','msp'):
            for method in ('Equal','Duration','Amount'):
                _,raw=make(source,method);check_package(self,raw)
                with ZipFile(BytesIO(raw)) as z:
                    for name in z.namelist():
                        if name.startswith('xl/charts/chart') and name.endswith('.xml'):check_line_fills(self,ET.fromstring(z.read(name)))

if __name__=='__main__':unittest.main()
