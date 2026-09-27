"""F2 MS-2 ordinary validation + PO milestone Policy A; no Excel PASS claim."""
from copy import deepcopy
from dataclasses import replace
from decimal import Decimal
from io import BytesIO
import json
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.worksheet._writer import ALL_TEMP_FILES
from app import app, MAX_BYTES, XLSX_TYPE, _conversion_lock
from model import normalize, parse_xml
from engine.amount import preview_amounts
from engine.domain import Settings, InputError, prepare, weight_bases
from engine.service import convert, inspect_amount
from test_f1 import activity_rows
from test_f1_chart_integrity import check_package, check_line_fills


def amount_xml(source='msp', values=('100', '300', '0', '500')):
    """Synthetic public test schedule only. Last activity is an explicit milestone."""
    root = ET.Element('Project' if source == 'msp' else 'APIBusinessObjects')
    project = root if source == 'msp' else ET.SubElement(root, 'Project')
    def add(parent, **fields):
        for k, v in fields.items(): ET.SubElement(parent, k).text = str(v)
    add(project, Name='F2 synthetic', **({'UID': 'p'} if source == 'msp' else {'ObjectId': 'p'}))
    definitions = ET.SubElement(project, 'ExtendedAttributes') if source == 'msp' else root
    for ident, name, dtype in [('11', 'Contract allocation', 'Number'), ('12', 'Other value', 'Cost'), ('13', 'Numeric looking text', 'Text')]:
        if source == 'msp':
            add(ET.SubElement(definitions, 'ExtendedAttribute'), FieldID=ident, Alias=name,
                FieldName=dtype+'1', CFType='5' if dtype == 'Number' else '7' if dtype == 'Text' else '9')
        else:
            add(ET.SubElement(definitions, 'UDFType'), ObjectId=ident, Title=name,
                SubjectArea='Activity', DataType='Double' if dtype == 'Number' else dtype)
    if source == 'msp':
        tasks = ET.SubElement(project, 'Tasks')
        add(ET.SubElement(tasks, 'Task'), UID='10', ID='0', Name='Structure', Summary='1', OutlineLevel='1', WBS='1')
    else:
        add(ET.SubElement(project, 'WBS'), ObjectId='10', Code='1', Name='Structure')
        tasks = project
    for i, val in enumerate(values, 1):
        task = ET.SubElement(tasks, 'Task' if source == 'msp' else 'Activity')
        if source == 'msp':
            add(task, UID=str(i), ID=str(i), Name=f'Work {i}', Summary='0', OutlineLevel='2',
                Start='2026-01-23', Finish='2026-02-13', Duration='PT8H', Milestone='1' if i == 4 else '0')
        else:
            add(task, ObjectId=str(i), Id=str(i), Name=f'Work {i}', WBSObjectId='10',
                PlannedStartDate='2026-01-23', PlannedFinishDate='2026-02-13', PlannedDuration='8',
                Type='Finish Milestone' if i == 4 else 'Task Dependent')
        for ident, raw in [('11', val), ('12', str(i*10)), ('13', '999')]:
            if raw is None: continue
            if source == 'msp': add(ET.SubElement(task, 'ExtendedAttribute'), FieldID=ident, Value=raw)
            else: add(ET.SubElement(task, 'UDF'), TypeObjectId=ident,
                      **{('CostValue' if ident == '12' else 'DoubleValue'): raw})
    root.set('xmlns', 'http://schemas.microsoft.com/project' if source == 'msp' else 'urn:p6')
    return ET.tostring(root)


def selected(source): return 'msp:field:11' if source == 'msp' else 'p6:udf:11'


def schedule(source='msp', values=('100','300','0','500')):
    raw = amount_xml(source, values)
    return normalize(parse_xml(raw), raw, production=True)


class AmountDomainTests(unittest.TestCase):
    def test_discovery_typed_identity_and_no_automatic_selection(self):
        for source in ('msp', 'p6'):
            s = schedule(source)
            self.assertEqual(len(s.amount_fields), 2)
            self.assertEqual({f.name for f in s.amount_fields}, {'Contract allocation','Other value'})
            self.assertIsNone(inspect_amount(amount_xml(source))['preview'])
            for field in (None, 'unknown', selected(source).replace('11','13')):
                with self.assertRaises(InputError): preview_amounts(s, field)
            with self.assertRaises(InputError): Settings('Amount')
            with self.assertRaises(InputError): Settings('Equal', amount_field=selected(source))

    def test_denominator_and_milestone_money_are_separate(self):
        for source in ('msp','p6'):
            s = schedule(source); p = preview_amounts(s, selected(source)).require_valid()
            self.assertEqual(p.total, Decimal(400)); self.assertEqual(p.bases, [100,300,0,0])
            self.assertEqual(p.rows[-1].value, Decimal(500))
            self.assertEqual([v/sum(p.bases) for v in p.bases],[.25,.75,0,0])
            self.assertEqual((p.positive_count,p.zero_count),(2,1))
            self.assertEqual(weight_bases(s, 'Amount', selected(source)),p.bases)
            self.assertEqual(weight_bases(s, 'Equal'),[1,1,1,0])
            self.assertEqual(weight_bases(s, 'Duration'),[8,8,8,0])

    def test_ordinary_missing_blank_and_invalid_reject(self):
        for source in ('msp','p6'):
            for bad in (None,'',' ','bad','-1','NaN','Infinity','1e309','1e-999','1_000','1,000'):
                with self.subTest(source=source,bad=bad):
                    s=schedule(source,(bad,'300','0','500'));p=preview_amounts(s,selected(source))
                    self.assertFalse(p.valid)
                    self.assertIn('1 (Work 1)',p.errors[0]);self.assertIn(selected(source),p.errors[0])
                    with self.assertRaises(InputError): p.require_valid()
                    self.assertEqual(weight_bases(s,'Equal'),[1,1,1,0])
                    self.assertEqual(weight_bases(s,'Duration'),[8,8,8,0])

    def test_policy_a_milestone_missing_blank_zero_positive(self):
        for source in ('msp','p6'):
            for raw in (None,'',' ','0','500.123456789'):
                p=preview_amounts(schedule(source,('100','300','0',raw)),selected(source)).require_valid()
                self.assertEqual(p.bases[-1],0);self.assertEqual(p.total,400)
                if raw is None or not raw.strip():
                    self.assertIsNone(p.rows[-1].value);self.assertEqual(len(p.warnings),1)
                else:
                    self.assertEqual(p.rows[-1].value,Decimal(raw));self.assertFalse(p.warnings)

    def test_policy_a_invalid_milestone_blocks(self):
        for source in ('msp','p6'):
            for bad in ('bad','-1','NaN','Infinity','1e309','1e-999','1_000'):
                with self.subTest(source=source,bad=bad):
                    p=preview_amounts(schedule(source,('100','300','0',bad)),selected(source))
                    self.assertFalse(p.valid);self.assertIn('4 (Work 4)',p.errors[0])
                    self.assertIn(selected(source),p.errors[0])
                    with self.assertRaises(InputError): p.bases

    def test_duplicate_activity_values_and_definitions_reject(self):
        for source in ('msp','p6'):
            s=schedule(source)
            for idx in (0,3):
                acts=list(s.activities);acts[idx]=replace(acts[idx],amount_field_values=acts[idx].amount_field_values+((selected(source),'50'),))
                p=preview_amounts(replace(s,activities=tuple(acts)),selected(source))
                self.assertIn('duplicate',p.errors[0])
            with self.assertRaisesRegex(InputError,'ambiguously'):
                preview_amounts(replace(s,amount_fields=s.amount_fields+(s.amount_fields[0],)),selected(source))

    def test_duplicate_xml_value_children_detected(self):
        for source in ('msp','p6'):
            root=parse_xml(amount_xml(source));tag='Value' if source=='msp' else 'DoubleValue'
            target=next(n for n in root.iter() if n.find(tag) is not None)
            ET.SubElement(target,tag).text='777'
            raw=ET.tostring(root);s=normalize(root,raw,production=True)
            self.assertIn('duplicate',preview_amounts(s,selected(source)).errors[0])

    def test_zero_total_and_overflow_total(self):
        for vals in [('0','0','0','999'),('1e308','1e308','0','500')]:
            p=preview_amounts(schedule(values=vals),selected('msp'))
            self.assertFalse(p.valid);self.assertIn('total',p.errors[0])
        s=schedule();s=replace(s,activities=(s.activities[-1],))
        self.assertFalse(preview_amounts(s,selected('msp')).valid)

    def test_precision_without_cent_rounding_and_zero_exponent(self):
        p=preview_amounts(schedule(values=('0.00001','0.00002','0e-999999999','500')),selected('msp'))
        self.assertTrue(p.valid);self.assertEqual(p.total,Decimal('0.00003'))
        self.assertEqual(p.bases,[.00001,.00002,0,0])

    def test_lookup_only_native_cost_and_undeclared_not_inferred(self):
        root=parse_xml(amount_xml());task=root.find('Tasks/Task[Summary="0"]')
        item=task.find('ExtendedAttribute');item.remove(item.find('Value'));ET.SubElement(item,'ValueGUID').text='lookup'
        ET.SubElement(task,'Cost').text='100000'
        raw=ET.tostring(root);p=preview_amounts(normalize(root,raw,production=True),selected('msp'))
        self.assertFalse(p.valid)
        root.remove(root.find('ExtendedAttributes'));raw=ET.tostring(root)
        self.assertFalse(normalize(root,raw,production=True).amount_fields)

    def test_duplicate_display_names_and_existing_identity_preserved(self):
        s=schedule();fields=tuple(replace(f,name='Same') for f in s.amount_fields)
        p=preview_amounts(replace(s,amount_fields=fields),'msp:field:12')
        self.assertEqual(p.total,60)
        raw=amount_xml();root=parse_xml(raw)
        before=normalize(root,raw,import_id='00000000-0000-0000-0000-000000000001',production=True)
        root.remove(root.find('ExtendedAttributes'))
        after=normalize(root,raw,import_id=before.import_id,production=True)
        self.assertEqual([(a.key,a.source_id,a.source_object_id,a.wbs) for a in before.activities],
                         [(a.key,a.source_id,a.source_object_id,a.wbs) for a in after.activities])

    def test_p6_integer_cost_and_nonactivity_declarations(self):
        root=parse_xml(amount_xml('p6'))
        root.find('UDFType/DataType').text='Integer'
        for node in root.iter('DoubleValue'): node.tag='IntegerValue'
        nonactivity=deepcopy(root.find('UDFType'));nonactivity.find('ObjectId').text='99'
        nonactivity.find('SubjectArea').text='Project';root.append(nonactivity)
        raw=ET.tostring(root);s=normalize(root,raw,production=True)
        self.assertEqual(len(s.amount_fields),2)
        self.assertEqual(preview_amounts(s,'p6:udf:11').bases,[100,300,0,0])
        self.assertEqual(preview_amounts(s,'p6:udf:12').total,60)


class AmountWorkbookTests(unittest.TestCase):
    def test_values_metadata_formula_and_chart_integrity(self):
        for source in ('msp','p6'):
            blob,info=convert(amount_xml(source),'Amount',amount_field=selected(source))
            self.assertEqual(info['warnings'],0);check_package(self,blob)
            with ZipFile(BytesIO(blob)) as z:
                for name in z.namelist():
                    if name.startswith('xl/charts/') and name.endswith('.xml'):
                        check_line_fills(self,ET.fromstring(z.read(name)))
            w=load_workbook(BytesIO(blob));ws=w['Main'];rows=activity_rows(ws)
            self.assertEqual([ws.cell(r,9).value for r in rows],[100,300,0,0])
            for r in rows:
                self.assertEqual(ws.cell(r,10).value,f'=I{r}/$I$5')
                self.assertIsNone(ws.cell(r+1,14).value)
            amount=w['Activity Amount'];arows=[r for r in range(4,amount.max_row+1) if amount.cell(r,7).value]
            self.assertEqual([amount.cell(r,4).value for r in arows],[100,300,0,500])
            self.assertTrue(all(not amount.cell(r,4).protection.locked for r in arows))
            self.assertEqual([amount.cell(r,5).value for r in arows],[ws.cell(r,12).value for r in rows])
            meta={r[0].value:r[1].value for r in w['_Metadata']}
            self.assertEqual(meta['amount_field_identity'],selected(source))
            self.assertEqual(meta['weight_method'],'Amount');self.assertEqual(meta['ordinary_amount_basis_total'],'400')
            self.assertEqual(meta['schema'],'bluebird-f2-1')
            self.assertEqual(len(meta['source_hash']),64)
            raw_meta=[json.loads(r[1].value) for r in w['_Metadata'] if r[0].value=='creation_amount']
            self.assertEqual(raw_meta[-1]['value'],'500')
            for sheet in ('Main','Monthly','Dashboard','Dashboard_Data'):
                for row in w[sheet]:
                    for c in row:
                        if c.data_type=='f': self.assertNotIn('Activity Amount',c.value)
            w.close()

    def test_policy_a_blank_zero_retention_warning_and_roundtrip(self):
        for raw in (None,'','0','500.5'):
            blob,info=convert(amount_xml(values=('100','300','0',raw)),'Amount',amount_field=selected('msp'))
            w=load_workbook(BytesIO(blob));ws=w['Activity Amount']
            r=max(r for r in range(4,ws.max_row+1) if ws.cell(r,7).value)
            self.assertEqual(ws.cell(r,4).value, None if raw in (None,'') else float(raw))
            self.assertEqual(info['warnings'],1 if raw in (None,'') else 0)
            if raw in (None,''):self.assertTrue(any('Missing milestone' in str(row[0].value) for row in w['Guide']))
            out=BytesIO();w.save(out);w.close();w=load_workbook(BytesIO(out.getvalue()))
            self.assertEqual(w[ws.title].cell(r,4).value,None if raw in (None,'') else float(raw));w.close()

    def test_equal_duration_do_not_import_or_validate_unused_money(self):
        for method in ('Equal','Duration'):
            blob,_=convert(amount_xml(values=('bad','-1','NaN','bad')),method)
            w=load_workbook(BytesIO(blob));ws=w['Activity Amount']
            self.assertTrue(all(ws.cell(r,4).value is None for r in range(4,ws.max_row+1)))
            self.assertNotIn('amount_field_identity',{r[0].value for r in w['_Metadata']});w.close()

    def test_all_f1_formula_and_chart_parts_unchanged(self):
        raw=amount_xml();a,_=convert(raw,'Equal');b,_=convert(raw,'Amount',amount_field=selected('msp'))
        wa=load_workbook(BytesIO(a));wb=load_workbook(BytesIO(b))
        for name in ('Main','Monthly','Dashboard','Dashboard_Data','progress'):
            formulas=lambda ws:{c.coordinate:c.value for row in ws for c in row if c.data_type=='f'}
            self.assertEqual(formulas(wa[name]),formulas(wb[name]))
        with ZipFile(BytesIO(a)) as za,ZipFile(BytesIO(b)) as zb:
            for name in za.namelist():
                if name.startswith(('xl/charts/','xl/drawings/')):self.assertEqual(za.read(name),zb.read(name),name)
        wa.close();wb.close()

    def test_duplicate_display_activity_ids_do_not_cross_assign_money(self):
        root=parse_xml(amount_xml());tasks=root.findall('Tasks/Task[Summary="0"]')
        for task in tasks:task.find('ID').text='SAME'
        blob,_=convert(ET.tostring(root),'Amount',amount_field='msp:field:11')
        w=load_workbook(BytesIO(blob));ws=w['Activity Amount']
        rows=[r for r in range(4,ws.max_row+1) if ws.cell(r,7).value]
        self.assertEqual([ws.cell(r,4).value for r in rows],[100,300,0,500])
        self.assertEqual(len({ws.cell(r,5).value for r in rows}),4);w.close()


class AmountRuntimeTests(unittest.TestCase):
    def setUp(self): self.client=TestClient(app)
    def post(self, path, raw, query=''):
        return self.client.post('/api/progress/'+path+'?'+query,content=raw,headers={'Content-Type':'application/xml'})

    def test_discover_validate_generate_both_sources(self):
        for source in ('msp','p6'):
            raw=amount_xml(source);before=set(ALL_TEMP_FILES)
            r=self.post('amount-preview',raw);self.assertEqual(r.status_code,200)
            self.assertIsNone(r.json()['preview']);self.assertEqual(len(r.json()['fields']),2)
            p=self.post('amount-preview',raw,'amount_field='+selected(source)).json()
            self.assertTrue(p['preview']['valid']);self.assertEqual(p['preview']['ordinary_total'],'400')
            q='method=Amount&amount_field='+selected(source)+'&source_hash='+p['source_hash']
            r=self.post('convert',raw,q);self.assertEqual(r.status_code,200,r.text[:100] if r.status_code!=200 else '')
            self.assertEqual(r.headers['content-type'],XLSX_TYPE);self.assertIn('progress.xlsx',r.headers['content-disposition'])
            self.assertEqual(r.headers['cache-control'],'no-store');check_package(self,r.content)
            self.assertEqual(set(ALL_TEMP_FILES),before)

    def test_server_validates_without_trusting_preview_and_rejects_stale_hash(self):
        q='method=Amount&amount_field=msp:field:11'
        r=self.post('convert',amount_xml(values=('100','300','0','bad')),q)
        self.assertEqual(r.status_code,422);self.assertIn('4 (Work 4)',r.json()['error']['message'])
        self.assertIn('msp:field:11',r.json()['error']['message'])
        self.assertEqual(self.post('convert',amount_xml(),q+'&source_hash=wrong').status_code,422)
        for query in ('method=Amount','method=Amount&amount_field=unknown',q+'&amount_field=msp:field:12'):
            self.assertEqual(self.post('convert',amount_xml(),query).status_code,422)

    def test_preview_reports_all_errors_and_warning_without_creating_xlsx(self):
        p=self.post('amount-preview',amount_xml(values=('bad','300','0',None)),'amount_field=msp:field:11')
        self.assertEqual(p.status_code,200);p=p.json()['preview']
        self.assertFalse(p['valid']);self.assertEqual(len(p['errors']),1);self.assertEqual(len(p['warnings']),1)
        self.assertIsNone(p['rows'][-1]['value'])

    def test_preview_bounded_safe_failures_and_busy(self):
        for raw,status in [(b'',400),(b'<bad',400),(b'<unknown/>',422),(b'x'*(MAX_BYTES+1),413),
            (b'<!DOCTYPE x [<!ENTITY y "bad">]><x>&y;</x>',422)]:
            self.assertEqual(self.post('amount-preview',raw).status_code,status)
        with patch('app.inspect_amount',side_effect=RuntimeError('PRIVATE')):
            r=self.post('amount-preview',amount_xml());self.assertEqual(r.status_code,500);self.assertNotIn('PRIVATE',r.text)
        with patch('app.monotonic',side_effect=[0,61]):
            self.assertEqual(self.post('amount-preview',amount_xml()).status_code,504)
        _conversion_lock.acquire()
        try:self.assertEqual(self.post('amount-preview',amount_xml()).status_code,503)
        finally:_conversion_lock.release()

    def test_preview_no_fields_and_duplicate_query(self):
        root=parse_xml(amount_xml());root.remove(root.find('ExtendedAttributes'));raw=ET.tostring(root)
        self.assertEqual(self.post('amount-preview',raw).json()['fields'],[])
        self.assertEqual(self.post('amount-preview',raw,'amount_field=x&amount_field=y').status_code,422)


if __name__=='__main__': unittest.main()
